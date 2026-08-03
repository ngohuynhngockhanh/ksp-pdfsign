from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
import subprocess
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .auth import CurrentUser, require_admin
from .config import get_settings
from .db import (JobRun, PayrollEmployee, PayrollEmployeeAlias, PayrollImport,
                 PayrollLine, PayrollPayment, PayrollPeriod, PayrollStatement,
                 PayrollWorkbookDraft, get_session)
from .payroll import (COL_MEAL, DEPENDENT_DEDUCTION, EMPLOYEE_INSURANCE_RATE,
                      EMPLOYER_INSURANCE_RATE, SELF_DEDUCTION, TRADE_UNION_RATE,
                      PayrollInput, apply_workbook_changes, calculate_annual_pit,
                      calculate_payroll,
                      insurance_base_cap, plan_annual_bonus, plan_net_target,
                      review_workbook)
from . import audit

router = APIRouter(prefix="/api/payroll", tags=["payroll"])
logger = logging.getLogger(__name__)


class EmployeeIn(BaseModel):
    code: str
    name: str
    position: str = ""
    base_salary: float = Field(ge=0)
    insurance_salary: float = Field(ge=0)
    meal_allowance: float = Field(default=0, ge=0)
    phone_allowance: float = Field(default=0, ge=0)
    fuel_allowance: float = Field(default=0, ge=0)
    responsibility_allowance: float = Field(default=0, ge=0)
    dependents: int = Field(default=0, ge=0)


class PeriodIn(BaseModel):
    month: str = Field(pattern=r"^20\d\d-(0[1-9]|1[0-2])$")


class LineIn(BaseModel):
    standard_days: float | None = Field(default=None, gt=0)
    actual_days: float | None = Field(default=None, ge=0)
    overtime_pay: float | None = Field(default=None, ge=0)
    bonus: float | None = Field(default=None, ge=0)
    other_taxable: float | None = Field(default=None, ge=0)
    unpaid_deduction: float | None = Field(default=None, ge=0)
    override_net: float | None = Field(default=None, ge=0)
    override_reason: str = ""

    @model_validator(mode="after")
    def require_reason(self):
        if self.override_net is not None and not self.override_reason.strip():
            raise ValueError("Giá trị ghi đè phải có lý do")
        return self


class WorkbookChange(BaseModel):
    row: int = Field(ge=15, le=200)
    meal_allowance: float | None = Field(default=None, ge=0)
    attendance_bonus: float | None = Field(default=None, ge=0)
    overtime_weekday_hours: float | None = Field(default=None, ge=0, le=400)
    overtime_weekend_hours: float | None = Field(default=None, ge=0, le=400)
    reason: str = Field(min_length=3, max_length=500)

    @model_validator(mode="after")
    def has_change(self):
        fields = (self.meal_allowance, self.attendance_bonus,
                  self.overtime_weekday_hours, self.overtime_weekend_hours)
        if all(value is None for value in fields):
            raise ValueError("Phải nhập ít nhất một khoản cần điều chỉnh")
        return self


class WorkbookDraftIn(BaseModel):
    changes: list[WorkbookChange] = Field(min_length=1, max_length=100)

    @model_validator(mode="after")
    def unique_rows(self):
        rows = [item.row for item in self.changes]
        if len(rows) != len(set(rows)):
            raise ValueError("Mỗi nhân viên chỉ được xuất hiện một lần trong bản nháp")
        return self


class NetTargetIn(BaseModel):
    row: int = Field(ge=15, le=200)
    target_net: float = Field(gt=0, le=1_000_000_000)
    available_weekday_ot_hours: float | None = Field(default=None, ge=0, le=400)
    available_weekend_ot_hours: float | None = Field(default=None, ge=0, le=400)
    allow_taxable_bonus: bool = False


class PaymentIn(BaseModel):
    employee_id: int = Field(gt=0)
    month: str = Field(pattern=r"^20\d\d-(0[1-9]|1[0-2])$")
    amount: float = Field(gt=0, le=1_000_000_000)
    status: str = Field(pattern=r"^(prepared|completed)$")
    paid_at: datetime | None = None
    bank_name: str = Field(default="", max_length=150)
    transaction_ref: str = Field(default="", max_length=150)
    note: str = Field(default="", max_length=500)


class CancelPaymentIn(BaseModel):
    reason: str = Field(min_length=3, max_length=500)


def _rclone_binary() -> str:
    """Duong dan rclone: uu tien PATH, sau do ~/.local/bin (systemd khong co PATH day du)."""
    local_rclone = Path.home() / ".local" / "bin" / "rclone"
    rclone = shutil.which("rclone") or (str(local_rclone) if local_rclone.is_file() else "")
    if not rclone:
        raise HTTPException(503, "Không tìm thấy rclone trên máy chủ")
    return rclone


def _drive_flags(settings) -> list[str]:
    return ["--drive-root-folder-id", settings.payroll_drive_folder_id,
            "--bind", settings.payroll_rclone_bind]


def _loads(value: str, default):
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError) as exc:
        # JSON hong trong DB -> tra default nhung PHAI log de ke toan biet du lieu
        # (override luong / snapshot / findings) co the bi mat, tranh sai lam lang.
        if value not in (None, "", "null"):
            logger.warning("payroll: bo qua JSON hong (%s): %.80r", type(exc).__name__, value)
        return default


def _normalized_employee_name(value: str) -> str:
    folded = unicodedata.normalize("NFKD", value.strip().lower())
    ascii_name = "".join(char for char in folded if not unicodedata.combining(char))
    return " ".join(ascii_name.replace("đ", "d").split())


def _grid_number(row: list[str], index: int) -> float:
    if index >= len(row):
        return 0.0
    text = str(row[index]).strip().replace(" ", "")
    if not text:
        return 0.0
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return 0.0


def _is_payroll_section_or_total(name: str) -> bool:
    normalized = _normalized_employee_name(name)
    return (normalized in {"chinh thuc", "thu viec"}
            or normalized.startswith("tong ") or normalized.startswith("cong "))


def _employee_for_excel_name(db: Session, name: str, position: str) -> PayrollEmployee:
    normalized = _normalized_employee_name(name)
    alias = db.scalar(select(PayrollEmployeeAlias).where(
        PayrollEmployeeAlias.normalized_name == normalized))
    if alias:
        employee = db.get(PayrollEmployee, alias.employee_id)
        if employee:
            return employee
    digest = hashlib.sha256(normalized.encode()).hexdigest()[:12].upper()
    code = f"HR-{digest}"
    employee = db.scalar(select(PayrollEmployee).where(PayrollEmployee.code == code))
    if not employee:
        employee = PayrollEmployee(code=code, name=" ".join(name.split()),
                                   position=position.strip())
        db.add(employee); db.flush()
    db.add(PayrollEmployeeAlias(normalized_name=normalized, employee_id=employee.id,
                                display_name=name.strip()))
    db.flush()
    return employee


def _latest_imports_for_year(db: Session, year: int) -> list[PayrollImport]:
    rows = db.scalars(select(PayrollImport).where(
        PayrollImport.month.like(f"{year:04d}-%")).order_by(
            PayrollImport.imported_at, PayrollImport.id)).all()
    latest: dict[str, PayrollImport] = {}
    for row in rows:
        filename = row.filename.lower()
        if "bản gốc trước cập nhật" in filename or "ban goc truoc cap nhat" in filename:
            continue
        latest[row.month] = row
    return list(latest.values())


def _rebuild_hr_statements(db: Session, year: int) -> None:
    imports = _latest_imports_for_year(db, year)
    db.query(PayrollStatement).filter(PayrollStatement.month.like(f"{year:04d}-%")).update(
        {PayrollStatement.is_current: False}, synchronize_session=False)
    for imported in imports:
        snapshot = _loads(imported.snapshot, {})
        for excel_row, row in enumerate(snapshot.get("grid", [])[14:], 15):
            name = str(row[1]).strip() if len(row) > 1 else ""
            if not name or _is_payroll_section_or_total(name):
                continue
            position = str(row[2]).strip() if len(row) > 2 else ""
            employee = _employee_for_excel_name(db, name, position)
            statement = db.scalar(select(PayrollStatement).where(
                PayrollStatement.employee_id == employee.id,
                PayrollStatement.month == imported.month,
                PayrollStatement.source_import_id == imported.id,
            ))
            values = {
                "source_row": excel_row,
                "is_current": True,
                "gross_income": _grid_number(row, 16),
                "employee_insurance": _grid_number(row, 26),
                "taxable_income_before_deductions": _grid_number(row, 30),
                "pit_withheld": _grid_number(row, 32),
                "net_payable": _grid_number(row, 35),
                "employer_cost": _grid_number(row, 36),
                "dependent_count": int(_grid_number(row, 28)),
            }
            if statement:
                for key, value in values.items():
                    setattr(statement, key, value)
            else:
                db.add(PayrollStatement(employee_id=employee.id, month=imported.month,
                                        source_import_id=imported.id, **values))
    db.commit()


def _payment_out(row: PayrollPayment) -> dict:
    return {
        "id": row.id, "employee_id": row.employee_id, "month": row.month,
        "amount": row.amount, "status": row.status,
        "paid_at": row.paid_at.isoformat() if row.paid_at else "",
        "bank_name": row.bank_name, "transaction_ref": row.transaction_ref,
        "note": row.note, "evidence_name": row.evidence_name,
        "evidence_missing": row.status == "completed" and not row.evidence_path,
        "cancel_reason": row.cancel_reason,
    }


def _employee(row: PayrollEmployee):
    return {key: getattr(row, key) for key in (
        "id", "code", "name", "position", "base_salary", "insurance_salary",
        "meal_allowance", "phone_allowance", "fuel_allowance",
        "responsibility_allowance", "dependents", "active",
    )}


def _recalculate(period: PayrollPeriod, line: PayrollLine):
    employee = _loads(line.employee_snapshot, {})
    month = date.fromisoformat(period.month + "-01")
    result = calculate_payroll(PayrollInput(
        month=month, base_salary=employee["base_salary"],
        insurance_salary=employee["insurance_salary"], standard_days=line.standard_days,
        actual_days=line.actual_days, meal_allowance=employee.get("meal_allowance", 0),
        phone_allowance=employee.get("phone_allowance", 0), fuel_allowance=employee.get("fuel_allowance", 0),
        responsibility_allowance=employee.get("responsibility_allowance", 0),
        dependents=employee.get("dependents", 0), overtime_pay=line.overtime_pay,
        bonus=line.bonus, other_taxable=line.other_taxable, unpaid_deduction=line.unpaid_deduction,
    )).to_dict()
    overrides = _loads(line.overrides, {})
    if "net_income" in overrides:
        result["canonical_net_income"] = result["net_income"]
        result["net_income"] = overrides["net_income"]
    line.computed = json.dumps(result)
    return result


def _period(db: Session, row: PayrollPeriod):
    lines = db.scalars(select(PayrollLine).where(PayrollLine.period_id == row.id).order_by(PayrollLine.id)).all()
    return {"id": row.id, "month": row.month, "version": row.version, "status": row.status,
            "findings": _loads(row.findings, []), "lines": [
                {"id": line.id, "employee_id": line.employee_id,
                 "employee": _loads(line.employee_snapshot, {}), "standard_days": line.standard_days,
                 "actual_days": line.actual_days, "overtime_pay": line.overtime_pay, "bonus": line.bonus,
                 "other_taxable": line.other_taxable, "unpaid_deduction": line.unpaid_deduction,
                 "computed": _loads(line.computed, {}), "overrides": _loads(line.overrides, {}),
                 "override_reason": line.override_reason} for line in lines]}


@router.get("/employees")
def list_employees(db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    return [_employee(x) for x in db.scalars(select(PayrollEmployee).order_by(PayrollEmployee.code)).all()]


@router.post("/employees")
def create_employee(payload: EmployeeIn, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    if db.scalar(select(PayrollEmployee).where(PayrollEmployee.code == payload.code.strip())):
        raise HTTPException(409, "Mã nhân viên đã tồn tại")
    values = payload.model_dump()
    values.update(code=payload.code.strip(), name=payload.name.strip())
    row = PayrollEmployee(**values)
    db.add(row); db.commit(); db.refresh(row)
    return _employee(row)


@router.get("/periods")
def list_periods(db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    rows = db.scalars(select(PayrollPeriod).order_by(PayrollPeriod.month.desc(), PayrollPeriod.version.desc())).all()
    return [_period(db, row) for row in rows]


@router.post("/periods")
def create_period(payload: PeriodIn, db: Session = Depends(get_session), user: CurrentUser = Depends(require_admin)):
    version = (db.scalar(select(func.max(PayrollPeriod.version)).where(PayrollPeriod.month == payload.month)) or 0) + 1
    row = PayrollPeriod(month=payload.month, version=version, created_by=user.id,
                        policy_snapshot=json.dumps({
                            "pit": "109/2025/QH15", "self": SELF_DEDUCTION,
                            "dependent": DEPENDENT_DEDUCTION, "kpcd": TRADE_UNION_RATE,
                            "insurance_cap": insurance_base_cap(date.fromisoformat(payload.month + "-01")),
                        }))
    db.add(row); db.flush()
    employees = db.scalars(select(PayrollEmployee).where(PayrollEmployee.active.is_(True))).all()
    for employee in employees:
        line = PayrollLine(period_id=row.id, employee_id=employee.id,
                           employee_snapshot=json.dumps(_employee(employee), ensure_ascii=False),
                           standard_days=22, actual_days=0, overtime_pay=0, bonus=0,
                           other_taxable=0, unpaid_deduction=0, overrides="{}")
        _recalculate(row, line); db.add(line)
    db.commit(); db.refresh(row)
    return _period(db, row)


def _get_period(db: Session, period_id: int):
    row = db.get(PayrollPeriod, period_id)
    if not row: raise HTTPException(404, "Không tìm thấy kỳ lương")
    return row


@router.put("/periods/{period_id}/lines/{line_id}")
def update_line(period_id: int, line_id: int, payload: LineIn, db: Session = Depends(get_session), user: CurrentUser = Depends(require_admin)):
    period = _get_period(db, period_id)
    if period.status == "locked": raise HTTPException(409, "Kỳ lương đã khóa")
    line = db.get(PayrollLine, line_id)
    if not line or line.period_id != period.id: raise HTTPException(404, "Không tìm thấy dòng lương")
    for field in ("standard_days", "actual_days", "overtime_pay", "bonus", "other_taxable", "unpaid_deduction"):
        value = getattr(payload, field)
        if value is not None: setattr(line, field, value)
    if payload.override_net is not None:
        line.overrides = json.dumps({"net_income": payload.override_net})
        line.override_reason = payload.override_reason.strip()
        audit.record(db, user.username, user.role, user.ip, "payroll_override",
                     f"period:{period.id}/line:{line.id}", line.override_reason)
    line.updated_at = datetime.now(timezone.utc); _recalculate(period, line)
    period.status = "draft"; db.commit()
    return _period(db, period)


@router.post("/periods/{period_id}/review")
def review_period(period_id: int, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    row = _get_period(db, period_id)
    if row.status == "locked": raise HTTPException(409, "Kỳ lương đã khóa")
    lines = db.scalars(select(PayrollLine).where(PayrollLine.period_id == row.id)).all()
    findings = []
    if not lines: findings.append({"level": "do", "message": "Kỳ lương chưa có nhân viên."})
    for line in lines:
        employee = _loads(line.employee_snapshot, {})
        stable = employee.get("responsibility_allowance", 0)
        if stable and employee.get("insurance_salary", 0) <= employee.get("base_salary", 0):
            findings.append({"level": "vang", "employee_code": employee.get("code"),
                             "message": "Kiểm tra phụ cấp trách nhiệm ổn định trong nền đóng BHXH."})
    row.findings = json.dumps(findings, ensure_ascii=False); row.status = "reviewed"
    row.reviewed_at = datetime.now(timezone.utc); db.commit()
    return _period(db, row)


@router.post("/periods/{period_id}/lock")
def lock_period(period_id: int, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    row = _get_period(db, period_id)
    if row.status != "reviewed": raise HTTPException(409, "Phải review trước khi khóa")
    if any(x.get("level") == "do" for x in _loads(row.findings, [])):
        raise HTTPException(409, "Vẫn còn lỗi mức độ đỏ")
    row.status = "locked"; row.locked_at = datetime.now(timezone.utc); db.commit()
    return _period(db, row)


@router.post("/imports")
async def upload_import(file: UploadFile = File(...), db: Session = Depends(get_session), user: CurrentUser = Depends(require_admin)):
    data = await file.read()
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(400, "Chỉ chấp nhận tệp XLSX")
    root = get_settings().data_path / "payroll_imports"; root.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(data).hexdigest(); path = root / f"{digest}.xlsx"; path.write_bytes(data)
    snapshot = review_workbook(path)
    row = PayrollImport(month=snapshot["month"], filename=Path(file.filename).name, sha256=digest,
                        snapshot=json.dumps(snapshot, ensure_ascii=False),
                        findings=json.dumps(snapshot["findings"], ensure_ascii=False), imported_by=user.id)
    db.add(row); db.commit(); db.refresh(row)
    return {"id": row.id, "filename": row.filename, **snapshot}


@router.get("/imports")
def list_imports(db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    rows = db.scalars(select(PayrollImport).order_by(PayrollImport.month.desc(), PayrollImport.id.desc())).all()
    return [{"id": row.id, "month": row.month, "filename": row.filename,
             "findings": _loads(row.findings, []), "imported_at": row.imported_at.isoformat()}
            for row in rows]


@router.get("/imports/{import_id}")
def get_import(import_id: int, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    row = db.get(PayrollImport, import_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy tệp bảng lương")
    return {"id": row.id, "month": row.month, "filename": row.filename,
            "findings": _loads(row.findings, []), "snapshot": _loads(row.snapshot, {})}


@router.get("/hr-summary")
def hr_summary(year: int = 2026, db: Session = Depends(get_session),
               _: CurrentUser = Depends(require_admin)):
    if year < 2026 or year > 2100:
        raise HTTPException(400, "Năm tổng hợp không hợp lệ")
    _rebuild_hr_statements(db, year)
    statements = db.scalars(select(PayrollStatement).where(
        PayrollStatement.month.like(f"{year:04d}-%"),
        PayrollStatement.is_current.is_(True),
    ).order_by(PayrollStatement.month)).all()
    employee_ids = {row.employee_id for row in statements}
    employees = {employee_id: db.get(PayrollEmployee, employee_id)
                 for employee_id in employee_ids}
    payments = db.scalars(select(PayrollPayment).where(
        PayrollPayment.employee_id.in_(employee_ids) if employee_ids else PayrollPayment.id == -1,
        PayrollPayment.month.like(f"{year:04d}-%"),
    ).order_by(PayrollPayment.created_at)).all()
    payments_by_key: dict[tuple[int, str], list[PayrollPayment]] = {}
    for payment in payments:
        payments_by_key.setdefault((payment.employee_id, payment.month), []).append(payment)

    draft_rows: dict[int, set[int]] = {}
    for imported in _latest_imports_for_year(db, year):
        draft = db.scalar(select(PayrollWorkbookDraft).where(
            PayrollWorkbookDraft.import_id == imported.id,
            PayrollWorkbookDraft.status == "reviewed",
        ).order_by(PayrollWorkbookDraft.id.desc()))
        if draft:
            draft_rows[imported.id] = {int(item.get("row", 0))
                                       for item in _loads(draft.changes, [])}

    grouped: dict[int, list[PayrollStatement]] = {}
    for statement in statements:
        grouped.setdefault(statement.employee_id, []).append(statement)
    result = []
    for employee_id, rows in grouped.items():
        employee = employees[employee_id]
        months = []
        for statement in rows:
            month_payments = payments_by_key.get((employee_id, statement.month), [])
            completed = [item for item in month_payments if item.status == "completed"]
            paid = sum(item.amount for item in completed)
            outstanding = max(statement.net_payable - paid, 0)
            pending_revision = statement.source_row in draft_rows.get(statement.source_import_id, set())
            missing_evidence = any(not item.evidence_path for item in completed)
            if pending_revision:
                reconciliation = "pending_revision"
            elif missing_evidence:
                reconciliation = "missing_evidence"
            elif outstanding == 0 and statement.net_payable > 0:
                reconciliation = "paid"
            elif paid > 0:
                reconciliation = "partial"
            else:
                reconciliation = "unpaid"
            months.append({
                "month": statement.month, "statement_id": statement.id,
                "source_import_id": statement.source_import_id,
                "gross_income": statement.gross_income,
                "employee_insurance": statement.employee_insurance,
                "employer_insurance": round(
                    statement.employee_insurance / EMPLOYEE_INSURANCE_RATE
                    * EMPLOYER_INSURANCE_RATE
                ) if statement.employee_insurance else 0,
                "pit_withheld": statement.pit_withheld,
                "net_payable": statement.net_payable, "paid": paid,
                "outstanding": outstanding, "reconciliation_status": reconciliation,
                "payments": [_payment_out(item) for item in month_payments],
            })
        taxable_before = sum(row.taxable_income_before_deductions for row in rows)
        employee_insurance = sum(row.employee_insurance for row in rows)
        dependent_months = sum(row.dependent_count for row in rows)
        annual = calculate_annual_pit(
            taxable_income_before_deductions=taxable_before,
            employee_insurance=employee_insurance,
            dependent_months=dependent_months,
        )
        withheld = sum(row.pit_withheld for row in rows)
        net_payable = sum(row.net_payable for row in rows)
        paid = sum(month["paid"] for month in months)
        result.append({
            "employee_id": employee_id, "code": employee.code, "name": employee.name,
            "position": employee.position, "net_payable": net_payable, "paid": paid,
            "outstanding": max(net_payable - paid, 0),
            "gross_income": sum(row.gross_income for row in rows),
            "tax": {"withheld": withheld, **annual,
                    "balance": annual["annual_pit"] - withheld,
                    "basis": "Dữ liệu thu nhập do công ty quản lý"},
            "months": months,
        })
    result.sort(key=lambda item: item["name"])
    return {
        "year": year, "employees": result,
        "totals": {
            "net_payable": sum(item["net_payable"] for item in result),
            "paid": sum(item["paid"] for item in result),
            "outstanding": sum(item["outstanding"] for item in result),
            "pit_withheld": sum(item["tax"]["withheld"] for item in result),
            "annual_pit": sum(item["tax"]["annual_pit"] for item in result),
        },
    }


@router.get("/forecast")
def payroll_forecast(year: int = 2026, growth_rate: float = 0,
                     db: Session = Depends(get_session),
                     _: CurrentUser = Depends(require_admin)):
    if year < 2026 or year > 2100 or growth_rate < -.2 or growth_rate > .5:
        raise HTTPException(400, "Tham số dự báo không hợp lệ")
    _rebuild_hr_statements(db, year)
    statements = db.scalars(select(PayrollStatement).where(
        PayrollStatement.month.like(f"{year:04d}-%"),
        PayrollStatement.is_current.is_(True),
    ).order_by(PayrollStatement.month)).all()
    if not statements:
        return {"year": year, "actual_through": "", "growth_rate": growth_rate,
                "future_months": [], "baseline": {"gross_income": 0, "net_payable": 0,
                "annual_pit": 0, "pit_withheld": 0}, "recommended": {"gross_bonus": 0,
                "additional_pit": 0, "net_bonus": 0, "effective_tax_rate": 0},
                "scenarios": [], "employees": [], "assumptions": []}

    grouped: dict[int, list[PayrollStatement]] = {}
    for statement in statements:
        grouped.setdefault(statement.employee_id, []).append(statement)
    actual_through = max(row.month for row in statements)
    latest_month = int(actual_through[-2:])
    future_numbers = list(range(latest_month + 1, 13))
    future_totals = {month: {"gross_income": 0.0, "net_payable": 0.0,
                             "pit_estimate": 0.0} for month in future_numbers}
    employee_rows = []
    employee_plans = []
    baseline_totals = {"gross_income": 0.0, "net_payable": 0.0,
                       "annual_pit": 0.0, "pit_withheld": 0.0}

    for employee_id, rows in grouped.items():
        rows = sorted(rows, key=lambda item: item.month)
        sample = rows[-3:]
        count = len(sample)
        average = lambda field: sum(float(getattr(item, field)) for item in sample) / count
        projected = []
        for step, month in enumerate(future_numbers, 1):
            factor = (1 + growth_rate) ** step
            values = {
                "month": f"{year:04d}-{month:02d}",
                "gross_income": round(average("gross_income") * factor),
                "net_payable": round(average("net_payable") * factor),
                "pit_estimate": round(average("pit_withheld") * factor),
                "taxable_income_before_deductions": round(
                    average("taxable_income_before_deductions") * factor),
                "employee_insurance": round(average("employee_insurance")),
                "dependent_count": round(average("dependent_count")),
            }
            projected.append(values)
            for key in ("gross_income", "net_payable", "pit_estimate"):
                future_totals[month][key] += values[key]

        taxable_before = sum(row.taxable_income_before_deductions for row in rows)
        taxable_before += sum(row["taxable_income_before_deductions"] for row in projected)
        employee_insurance = sum(row.employee_insurance for row in rows)
        employee_insurance += sum(row["employee_insurance"] for row in projected)
        dependent_months = sum(row.dependent_count for row in rows)
        dependent_months += sum(row["dependent_count"] for row in projected)
        reference_gross = round(average("gross_income"))
        plan = plan_annual_bonus(
            taxable_income_before_deductions=taxable_before,
            employee_insurance=employee_insurance,
            dependent_months=dependent_months,
            reference_monthly_gross=reference_gross,
        )
        annual = calculate_annual_pit(
            taxable_income_before_deductions=taxable_before,
            employee_insurance=employee_insurance,
            dependent_months=dependent_months,
        )
        projected_gross = sum(row["gross_income"] for row in projected)
        projected_net = sum(row["net_payable"] for row in projected)
        projected_pit = sum(row["pit_estimate"] for row in projected)
        baseline_totals["gross_income"] += sum(row.gross_income for row in rows) + projected_gross
        baseline_totals["net_payable"] += sum(row.net_payable for row in rows) + projected_net
        baseline_totals["annual_pit"] += annual["annual_pit"]
        baseline_totals["pit_withheld"] += sum(row.pit_withheld for row in rows) + projected_pit
        employee = db.get(PayrollEmployee, employee_id)
        recommended = plan["recommended"]
        employee_rows.append({
            "employee_id": employee_id, "name": employee.name if employee else f"#{employee_id}",
            "reference_monthly_gross": reference_gross,
            "next_band_headroom": plan["next_band_headroom"],
            "recommended_bonus": recommended["gross_bonus"],
            "additional_pit": recommended["additional_pit"],
            "net_bonus": recommended["net_bonus"],
        })
        employee_plans.append(plan)

    scenarios = []
    for index, key in enumerate(("none", "half_month", "one_month", "one_half_month")):
        parts = [plan["scenarios"][index] for plan in employee_plans]
        gross = round(sum(part["gross_bonus"] for part in parts))
        tax = round(sum(part["additional_pit"] for part in parts))
        scenarios.append({
            "key": key, "label": parts[0]["label"] if parts else key,
            "gross_bonus": gross, "additional_pit": tax,
            "net_bonus": max(gross - tax, 0),
            "effective_tax_rate": round(tax / gross, 4) if gross else 0,
        })
    recommended_parts = [plan["recommended"] for plan in employee_plans]
    recommended_gross = round(sum(part["gross_bonus"] for part in recommended_parts))
    recommended_tax = round(sum(part["additional_pit"] for part in recommended_parts))
    recommended = {
        "gross_bonus": recommended_gross, "additional_pit": recommended_tax,
        "net_bonus": max(recommended_gross - recommended_tax, 0),
        "effective_tax_rate": round(recommended_tax / recommended_gross, 4) if recommended_gross else 0,
    }
    return {
        "year": year, "actual_through": actual_through, "growth_rate": growth_rate,
        "future_months": [{"month": f"{year:04d}-{month:02d}",
                           **{key: round(value) for key, value in totals.items()}}
                          for month, totals in future_totals.items()],
        "baseline": {key: round(value) for key, value in baseline_totals.items()},
        "recommended": recommended, "scenarios": scenarios,
        "employees": employee_rows,
        "assumptions": [
            "Dự phóng theo bình quân tối đa 3 tháng gần nhất của từng nhân viên.",
            "Tăng trưởng được áp lũy tiến theo từng tháng còn lại; mặc định là 0%.",
            "Thưởng được coi là thu nhập chịu TNCN và không tự động làm tăng nền BHXH.",
            "Mức tham khảo ưu tiên không đẩy phần thu nhập sang bậc thuế năm kế tiếp.",
        ],
    }


@router.post("/payments")
def create_payment(payload: PaymentIn, db: Session = Depends(get_session),
                   user: CurrentUser = Depends(require_admin)):
    employee = db.get(PayrollEmployee, payload.employee_id)
    if not employee:
        raise HTTPException(404, "Không tìm thấy nhân viên")
    statement = db.scalar(select(PayrollStatement).where(
        PayrollStatement.employee_id == payload.employee_id,
        PayrollStatement.month == payload.month,
        PayrollStatement.is_current.is_(True),
    ))
    if not statement:
        raise HTTPException(409, "Chưa có bảng lương chính thức của nhân viên trong tháng này")
    draft = db.scalar(select(PayrollWorkbookDraft).where(
        PayrollWorkbookDraft.import_id == statement.source_import_id,
        PayrollWorkbookDraft.status == "reviewed",
    ).order_by(PayrollWorkbookDraft.id.desc()))
    if draft and statement.source_row in {int(item.get("row", 0)) for item in _loads(draft.changes, [])}:
        raise HTTPException(409, "Dòng lương đang chờ chốt bản điều chỉnh; chưa thể ghi thanh toán")
    paid = db.scalar(select(func.coalesce(func.sum(PayrollPayment.amount), 0)).where(
        PayrollPayment.employee_id == payload.employee_id,
        PayrollPayment.month == payload.month,
        PayrollPayment.status == "completed",
    )) or 0
    if payload.status == "completed" and paid + payload.amount > statement.net_payable:
        raise HTTPException(409, "Số tiền vượt quá phần lương còn phải trả")
    row = PayrollPayment(
        employee_id=payload.employee_id, month=payload.month, amount=payload.amount,
        status=payload.status,
        paid_at=payload.paid_at or (datetime.now(timezone.utc) if payload.status == "completed" else None),
        bank_name=payload.bank_name.strip(), transaction_ref=payload.transaction_ref.strip(),
        note=payload.note.strip(), created_by=user.id,
    )
    db.add(row); db.commit(); db.refresh(row)
    audit.record(db, user.username, user.role, user.ip, "payroll_payment_create",
                 f"payment:{row.id}", f"{row.month}; {int(row.amount)} đồng; {row.status}")
    return _payment_out(row)


@router.post("/payments/{payment_id}/evidence")
async def upload_payment_evidence(payment_id: int, file: UploadFile = File(...),
                                  db: Session = Depends(get_session),
                                  user: CurrentUser = Depends(require_admin)):
    payment = db.get(PayrollPayment, payment_id)
    if not payment or payment.status == "cancelled":
        raise HTTPException(404, "Không tìm thấy thanh toán đang hiệu lực")
    filename = Path(file.filename or "").name
    suffix = Path(filename).suffix.lower()
    allowed = {".pdf": "application/pdf", ".png": "image/png",
               ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
    if suffix not in allowed:
        raise HTTPException(400, "Chứng từ chỉ chấp nhận PDF, PNG hoặc JPG")
    if file.content_type and file.content_type != allowed[suffix]:
        raise HTTPException(400, "Loại nội dung chứng từ không khớp phần mở rộng")
    data = await file.read(10_000_001)
    if not data or len(data) > 10_000_000:
        raise HTTPException(400, "Chứng từ phải có dung lượng từ 1 byte đến 10 MB")
    signatures = {".pdf": (b"%PDF",), ".png": (b"\x89PNG\r\n\x1a\n",),
                  ".jpg": (b"\xff\xd8\xff",), ".jpeg": (b"\xff\xd8\xff",)}
    if not any(data.startswith(signature) for signature in signatures[suffix]):
        raise HTTPException(400, "Nội dung chứng từ không đúng định dạng tệp")
    digest = hashlib.sha256(data).hexdigest()
    root = get_settings().data_path / "payroll_evidence" / str(payment.id)
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{digest}{suffix}"
    path.write_bytes(data)
    payment.evidence_path = str(path); payment.evidence_name = filename
    payment.evidence_sha256 = digest; payment.updated_at = datetime.now(timezone.utc)
    db.commit(); db.refresh(payment)
    audit.record(db, user.username, user.role, user.ip, "payroll_payment_evidence",
                 f"payment:{payment.id}", filename)
    return _payment_out(payment)


@router.get("/payments/{payment_id}/evidence")
def download_payment_evidence(payment_id: int, db: Session = Depends(get_session),
                              _: CurrentUser = Depends(require_admin)):
    payment = db.get(PayrollPayment, payment_id)
    if not payment or not payment.evidence_path or not Path(payment.evidence_path).is_file():
        raise HTTPException(404, "Không tìm thấy chứng từ thanh toán")
    return FileResponse(payment.evidence_path, filename=payment.evidence_name)


@router.post("/payments/{payment_id}/cancel")
def cancel_payment(payment_id: int, payload: CancelPaymentIn,
                   db: Session = Depends(get_session),
                   user: CurrentUser = Depends(require_admin)):
    payment = db.get(PayrollPayment, payment_id)
    if not payment:
        raise HTTPException(404, "Không tìm thấy thanh toán")
    if payment.status == "cancelled":
        raise HTTPException(409, "Thanh toán đã được hủy trước đó")
    payment.status = "cancelled"; payment.cancel_reason = payload.reason.strip()
    payment.updated_at = datetime.now(timezone.utc)
    db.commit(); db.refresh(payment)
    audit.record(db, user.username, user.role, user.ip, "payroll_payment_cancel",
                 f"payment:{payment.id}", payment.cancel_reason)
    return _payment_out(payment)


def _import_source(row: PayrollImport) -> Path:
    settings = get_settings()
    drive_path = settings.data_path / "payroll_drive" / row.drive_file_id
    upload_path = settings.data_path / "payroll_imports" / f"{row.sha256}.xlsx"
    path = drive_path if row.drive_file_id and drive_path.is_file() else upload_path
    if not path.is_file():
        raise HTTPException(404, "Không tìm thấy tệp Excel gốc trên máy chủ")
    return path


def _formula_hours(value: object) -> float:
    # Cong thuc OT do he thong sinh: =(D/E/8)*1.5*<gio> hoac =(D/E/8)*2*<gio>.
    # Chi lay factor cuoi (so gio); KHONG nhan voi he so 1.5/2 keo theo.
    # Bat buoc dung cau truc *<rate>*<gio> de khong bat nham so trong (D/E/8).
    if not isinstance(value, str) or not value.startswith("="):
        return 0
    grouped = re.search(
        r"\*\s*[0-9]+(?:\.[0-9]+)?\s*\*\s*\(\s*([0-9]+(?:\.[0-9]+)?)\s*\*\s*([0-9]+(?:\.[0-9]+)?)\s*\)\s*$",
        value,
    )
    if grouped:
        return float(grouped.group(1)) * float(grouped.group(2))
    match = re.search(r"\*\s*[0-9]+(?:\.[0-9]+)?\s*\*\s*([0-9]+(?:\.[0-9]+)?)\s*\)?\s*$", value)
    if not match:
        return 0
    return float(match.group(1))


@router.post("/imports/{import_id}/net-target")
def import_net_target(import_id: int, payload: NetTargetIn,
                      db: Session = Depends(get_session),
                      _: CurrentUser = Depends(require_admin)):
    imported = db.get(PayrollImport, import_id)
    if not imported:
        raise HTTPException(404, "Không tìm thấy tệp bảng lương")
    source = _import_source(imported)
    # Khong dung read-only: mot so file cu khai bao sai dimension A1:A1000.
    values_book = load_workbook(source, data_only=True, read_only=False)
    formulas_book = load_workbook(source, data_only=False, read_only=False)
    values = values_book[values_book.sheetnames[0]]
    formulas = formulas_book[formulas_book.sheetnames[0]]
    row = payload.row
    if values.cell(row, 2).value in (None, "") and values.cell(row, 3).value in (None, ""):
        raise HTTPException(400, "Dòng đã chọn không có nhân viên")
    month_text = imported.month or review_workbook(source)["month"]
    if not month_text:
        raise HTTPException(400, "Không xác định được tháng của bảng lương")

    def number(column: int) -> float:
        value = values.cell(row, column).value
        return float(value) if isinstance(value, (int, float)) else 0

    insurance = sum(number(column) for column in range(24, 28))
    if insurance <= 0:
        insurance = number(18) * .105
    try:
        return plan_net_target(
            month=date.fromisoformat(month_text + "-01"), current_net=number(36),
            target_net=payload.target_net, current_pit=number(33),
            current_employee_insurance=insurance, base_salary=number(4),
            standard_days=number(5), current_meal_allowance=number(COL_MEAL),
            available_weekday_ot_hours=payload.available_weekday_ot_hours,
            available_weekend_ot_hours=payload.available_weekend_ot_hours,
            current_weekday_ot_hours=_formula_hours(formulas.cell(row, 14).value),
            current_weekend_ot_hours=_formula_hours(formulas.cell(row, 15).value),
            current_overtime_pay=number(14) + number(15), current_gross=number(17),
            current_employer_cost=number(37),
            allow_taxable_bonus=payload.allow_taxable_bonus,
            current_performance_bonus=number(13), pit_taxable_income=number(32),
            pit_zero_headroom=max(number(27) + number(28) + number(30) - number(31), 0),
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


def _draft_out(row: PayrollWorkbookDraft) -> dict:
    return {"id": row.id, "import_id": row.import_id, "status": row.status,
            "changes": _loads(row.changes, []), "findings": _loads(row.findings, []),
            "drive_filename": row.drive_filename, "updated_at": row.updated_at.isoformat()}


@router.get("/imports/{import_id}/draft")
def latest_draft(import_id: int, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    row = db.scalar(select(PayrollWorkbookDraft).where(PayrollWorkbookDraft.import_id == import_id)
                    .order_by(PayrollWorkbookDraft.id.desc()))
    return {"draft": _draft_out(row) if row else None}


@router.post("/imports/{import_id}/draft")
def save_draft(import_id: int, payload: WorkbookDraftIn, db: Session = Depends(get_session),
               user: CurrentUser = Depends(require_admin)):
    imported = db.get(PayrollImport, import_id)
    if not imported:
        raise HTTPException(404, "Không tìm thấy tệp bảng lương")
    source = _import_source(imported)
    change_values = [item.model_dump() for item in payload.changes]
    draft = PayrollWorkbookDraft(import_id=import_id, status="draft", created_by=user.id,
                                 changes=json.dumps(change_values, ensure_ascii=False),
                                 updated_at=datetime.now(timezone.utc))
    db.add(draft); db.flush()
    output = get_settings().data_path / "payroll_drafts" / f"draft-{draft.id}.xlsx"
    apply_workbook_changes(source, output, change_values)
    draft.local_path = str(output)
    db.commit(); db.refresh(draft)
    audit.record(db, user.username, user.role, user.ip, "payroll_draft_save", f"draft:{draft.id}",
                 f"{len(payload.changes)} dòng thay đổi")
    return _draft_out(draft)


@router.post("/drafts/{draft_id}/review")
def review_draft(draft_id: int, db: Session = Depends(get_session), user: CurrentUser = Depends(require_admin)):
    draft = db.get(PayrollWorkbookDraft, draft_id)
    if not draft or not Path(draft.local_path).is_file():
        raise HTTPException(404, "Không tìm thấy bản nháp bảng lương")
    if draft.status != "draft":
        raise HTTPException(409, "Chỉ bản nháp mới được chạy review")
    snapshot = review_workbook(Path(draft.local_path))
    draft.findings = json.dumps(snapshot["findings"], ensure_ascii=False)
    draft.status = "reviewed"
    draft.reviewed_at = datetime.now(timezone.utc); draft.updated_at = draft.reviewed_at
    db.commit(); db.refresh(draft)
    audit.record(db, user.username, user.role, user.ip, "payroll_draft_review", f"draft:{draft.id}",
                 f"{len(snapshot['findings'])} cảnh báo")
    return _draft_out(draft)


@router.post("/drafts/{draft_id}/upload-drive")
def upload_draft(draft_id: int, db: Session = Depends(get_session), user: CurrentUser = Depends(require_admin)):
    draft = db.get(PayrollWorkbookDraft, draft_id)
    if not draft or not Path(draft.local_path).is_file():
        raise HTTPException(404, "Không tìm thấy bản nháp bảng lương")
    if draft.status != "reviewed":
        raise HTTPException(409, "Phải chạy review trước khi gửi lên Google Drive")
    if any(item.get("level") == "do" for item in _loads(draft.findings, [])):
        raise HTTPException(409, "Bản nháp vẫn còn cảnh báo mức đỏ")
    imported = db.get(PayrollImport, draft.import_id)
    if not imported:
        raise HTTPException(404, "Không tìm thấy tệp bảng lương gốc")
    stem = Path(imported.filename).stem
    filename = imported.filename if imported.drive_file_id else f"{stem} - đã review - bản {draft.id}.xlsx"
    settings = get_settings()
    commands = []
    if imported.drive_file_id:
        backup = f"{stem} - bản gốc trước cập nhật - bản {draft.id}.xlsx"
        commands.append([_rclone_binary(), "copyto",
                         f"{settings.payroll_drive_remote}{imported.filename}",
                         f"{settings.payroll_drive_remote}{backup}", *_drive_flags(settings)])
    commands.append([_rclone_binary(), "copyto", draft.local_path,
                     f"{settings.payroll_drive_remote}{filename}", *_drive_flags(settings)])
    try:
        for index, command in enumerate(commands):
            result = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
            if result.returncode:
                action = "sao lưu file gốc" if index == 0 and len(commands) > 1 else "cập nhật file Drive"
                raise HTTPException(502, f"Không thể {action}: {result.stderr[-300:]}")
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(504, "Gửi bản nháp lên Drive vượt quá thời gian chờ") from exc
    draft.status = "uploaded"; draft.drive_filename = filename
    draft.uploaded_at = datetime.now(timezone.utc); draft.updated_at = draft.uploaded_at
    db.commit(); db.refresh(draft)
    audit.record(db, user.username, user.role, user.ip, "payroll_draft_upload", f"draft:{draft.id}", filename)
    return _draft_out(draft)


def _update_job(db: Session, job: JobRun, **stats) -> None:
    current = _loads(job.stats, {})
    current.update(stats)
    job.stats = json.dumps(current, ensure_ascii=False)
    db.commit()


def _run_drive_sync(job_id: int, user_id: int) -> None:
    gen = get_session()
    db = next(gen)
    job = db.get(JobRun, job_id)
    if not job:
        gen.close()
        return
    try:
        _update_job(db, job, phase="connecting", progress=10, message="Đang kết nối Google Drive")
        result = _sync_drive_files(db, user_id, job)
        job.status = "success"
        job.stats = json.dumps({**result, "phase": "done", "progress": 100,
                                "message": f"Hoàn tất {len(result['files'])} file"}, ensure_ascii=False)
    except Exception as exc:
        logger.exception("payroll: dong bo Drive that bai (job %s)", job_id)
        db.rollback()
        job = db.get(JobRun, job_id)
        if job is None:  # Job bi xoa giua chung -> khong con gi de cap nhat.
            gen.close()
            return
        job.status = "failed"
        job.error = str(exc)[:1000]
    job.finished_at = datetime.now(timezone.utc)
    db.commit()
    gen.close()


def _sync_drive_files(db: Session, user_id: int, job: JobRun | None = None) -> dict:
    settings = get_settings()
    target = settings.data_path / "payroll_drive"; target.mkdir(parents=True, exist_ok=True)
    command = [_rclone_binary(), "copy", settings.payroll_drive_remote, str(target),
               *_drive_flags(settings), "--include", "*.xlsx", "--max-depth", "1"]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(504, "Đồng bộ Drive vượt quá thời gian chờ") from exc
    if result.returncode: raise HTTPException(502, f"Đồng bộ Drive gặp lỗi: {result.stderr[-300:]}")
    if job:
        _update_job(db, job, phase="reviewing", progress=55, message="Đang kiểm tra các file Excel")
    imported = 0
    summaries = []
    paths = [path for path in sorted(target.glob("*.xlsx")) if not path.name.startswith("~$")]
    for index, path in enumerate(paths, 1):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        try:
            snapshot = review_workbook(path)
        except Exception as exc:  # File loi khong duoc lam hong ca dot sync.
            logger.warning("payroll: khong doc duoc %s: %s", path.name, exc)
            summaries.append({"filename": path.name, "findings": 0,
                              "error": f"Không đọc được tệp Excel: {type(exc).__name__}"})
            continue
        summaries.append({"filename": path.name, "findings": len(snapshot["findings"])})
        existing = db.scalar(select(PayrollImport).where(PayrollImport.sha256 == digest))
        if existing:
            existing.month = snapshot["month"]
            existing.filename = path.name
            existing.snapshot = json.dumps(snapshot, ensure_ascii=False)
            existing.findings = json.dumps(snapshot["findings"], ensure_ascii=False)
        else:
            db.add(PayrollImport(month=snapshot["month"], filename=path.name, drive_file_id=path.name,
                                 sha256=digest, snapshot=json.dumps(snapshot, ensure_ascii=False),
                                 findings=json.dumps(snapshot["findings"], ensure_ascii=False),
                                 imported_by=user_id))
            imported += 1
        if job:
            _update_job(db, job, phase="reviewing", progress=55 + round(index / max(len(paths), 1) * 40),
                        message=f"Đang kiểm tra file {index}/{len(paths)}")
    db.commit()
    return {"files": [x["filename"] for x in summaries], "summaries": summaries,
            "imported": imported, "read_only": True}


@router.post("/sync-drive")
def start_drive_sync(background: BackgroundTasks, db: Session = Depends(get_session),
                     user: CurrentUser = Depends(require_admin)):
    running = db.scalar(select(JobRun).where(JobRun.kind == "payroll_drive_sync",
                                               JobRun.status == "running").order_by(JobRun.id.desc()))
    if running:
        return {"job_id": running.id, "status": running.status}
    job = JobRun(kind="payroll_drive_sync", status="running",
                 stats=json.dumps({"phase": "queued", "progress": 2, "message": "Đã xếp hàng đồng bộ"}))
    db.add(job); db.commit(); db.refresh(job)
    background.add_task(_run_drive_sync, job.id, user.id)
    return {"job_id": job.id, "status": job.status}


@router.get("/sync-drive/status")
def drive_sync_status(db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    job = db.scalar(select(JobRun).where(JobRun.kind == "payroll_drive_sync").order_by(JobRun.id.desc()))
    if not job:
        return {"job": None}
    return {"job": {"id": job.id, "status": job.status, "stats": _loads(job.stats, {}),
                    "error": job.error, "started_at": job.started_at.isoformat(),
                    "finished_at": job.finished_at.isoformat() if job.finished_at else ""}}


@router.get("/periods/{period_id}/export")
def export_period(period_id: int, db: Session = Depends(get_session), _: CurrentUser = Depends(require_admin)):
    period = _get_period(db, period_id); data = _period(db, period)
    workbook = Workbook(); sheet = workbook.active; sheet.title = period.month
    sheet.append(["Mã NV", "Họ tên", "Lương cơ bản", "Ngày công", "Tổng thu nhập", "BH NV", "TNCN", "Thực lĩnh", "Chi phí DN"])
    for line in data["lines"]:
        emp, calc = line["employee"], line["computed"]
        sheet.append([emp["code"], emp["name"], emp["base_salary"], line["actual_days"], calc["gross_income"], calc["employee_insurance"], calc["pit"], calc["net_income"], calc["total_employer_cost"]])
    root = get_settings().data_path / "exports"; root.mkdir(parents=True, exist_ok=True)
    path = root / f"bang-luong-{period.month}-v{period.version}.xlsx"; workbook.save(path)
    return FileResponse(path, filename=path.name)
