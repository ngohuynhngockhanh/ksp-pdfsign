"""Tinh luong Viet Nam 2026 va review file Excel luong hien huu."""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path
import re
from itertools import islice
from typing import Any

from openpyxl import load_workbook


SELF_DEDUCTION = 15_500_000
DEPENDENT_DEDUCTION = 6_200_000

# Ty le dong bao hiem (NLD 10,5% = BHXH 8 + BHYT 1,5 + BHTN 1; DN 21,5%; KPCD 2%).
EMPLOYEE_INSURANCE_RATE = .105
EMPLOYER_INSURANCE_RATE = .215
TRADE_UNION_RATE = .02

# Tran luong lam can cu dong BHXH bat buoc = 20 x muc tham chieu (luong co so),
# theo diem d khoan 1 Dieu 31 Luat BHXH 2024. Luong co so doi giua nam 2026:
#   - Truoc 01/07/2026: 2.340.000 (ND 73/2024) -> tran 46,8 trieu
#   - Tu 01/07/2026:    2.530.000 (ND 161/2026) -> tran 50,6 trieu
# Luu y: BHTN thuc te co tran rieng (20 x luong toi thieu vung, cao hon nhieu); o day
# ap chung tran luong co so cho toan bo phan bao hiem de khong tinh VUOT tran BHXH/BHYT.
BASE_REFERENCE_BEFORE_JUL_2026 = 2_340_000
BASE_REFERENCE_FROM_JUL_2026 = 2_530_000
INSURANCE_CAP_MULTIPLIER = 20

# Cot Excel dung chung cho review/edit/planner (1-based). Cot tien an la I (9);
# review, apply_workbook_changes va import_net_target phai tham chieu cung mot cot.
COL_MEAL = 9
COL_MEAL_LETTER = "I"


def insurance_base_cap(month: date) -> float:
    """Tran luong lam can cu dong BHXH bat buoc tai thoi diem cua ky luong."""
    reference = (BASE_REFERENCE_FROM_JUL_2026 if month >= date(2026, 7, 1)
                 else BASE_REFERENCE_BEFORE_JUL_2026)
    return reference * INSURANCE_CAP_MULTIPLIER


@dataclass
class PayrollInput:
    month: date
    base_salary: float
    standard_days: float
    actual_days: float
    insurance_salary: float | None = None
    meal_allowance: float = 0
    phone_allowance: float = 0
    fuel_allowance: float = 0
    responsibility_allowance: float = 0
    overtime_pay: float = 0
    bonus: float = 0
    other_taxable: float = 0
    unpaid_deduction: float = 0
    dependents: int = 0


@dataclass
class PayrollResult:
    prorated_salary: float
    gross_income: float
    insurance_salary: float
    employee_insurance: float
    employer_insurance: float
    trade_union_fee: float
    taxable_income_before_deductions: float
    pit_taxable_income: float
    pit: float
    net_income: float
    total_employer_cost: float

    def to_dict(self) -> dict[str, float]:
        return asdict(self)


def _money(value: float) -> float:
    return float(round(max(value, 0)))


def _signed_money(value: float) -> float:
    return float(round(value))


def calculate_pit(taxable_income: float) -> float:
    """Bieu thue luy tien 5 bac ap dung thu nhap tien luong tu nam 2026."""
    remaining = max(taxable_income, 0)
    tax = 0.0
    previous = 0.0
    for ceiling, rate in ((10e6, .05), (30e6, .10), (60e6, .20), (100e6, .30)):
        portion = min(remaining, ceiling - previous)
        tax += portion * rate
        remaining -= portion
        previous = ceiling
        if remaining <= 0:
            return _money(tax)
    return _money(tax + remaining * .35)


def calculate_payroll(data: PayrollInput) -> PayrollResult:
    if data.standard_days <= 0 or data.actual_days < 0:
        raise ValueError("Ngày công không hợp lệ")
    salary = data.base_salary * min(data.actual_days, data.standard_days) / data.standard_days
    gross = salary + data.meal_allowance + data.phone_allowance + data.fuel_allowance
    gross += data.responsibility_allowance + data.overtime_pay + data.bonus + data.other_taxable
    insurance_base = data.insurance_salary if data.insurance_salary is not None else data.base_salary
    # Ap tran 20 x muc tham chieu: phan luong vuot tran khong lam can cu dong BHXH.
    insurance_base = min(insurance_base, insurance_base_cap(data.month))
    employee_insurance = insurance_base * EMPLOYEE_INSURANCE_RATE
    employer_insurance = insurance_base * EMPLOYER_INSURANCE_RATE
    trade_union_fee = insurance_base * TRADE_UNION_RATE
    meal_exempt_cap = 1_200_000 if data.month >= date(2026, 7, 1) else 730_000
    taxable_before = gross - min(data.meal_allowance, meal_exempt_cap) - data.overtime_pay
    # Dien thoai/xang chi duoc mien khi co quy che/chung tu; mac dinh bao thu de review.
    pit_base = taxable_before - employee_insurance - SELF_DEDUCTION
    pit_base -= max(data.dependents, 0) * DEPENDENT_DEDUCTION
    pit = calculate_pit(pit_base)
    net = gross - employee_insurance - pit - data.unpaid_deduction
    employer_cost = gross + employer_insurance + trade_union_fee
    return PayrollResult(*map(_money, (
        salary, gross, insurance_base, employee_insurance, employer_insurance,
        trade_union_fee, taxable_before, pit_base, pit, net, employer_cost,
    )))


def plan_net_target(*, month: date, current_net: float, target_net: float,
                    current_pit: float, current_employee_insurance: float,
                    base_salary: float, standard_days: float, current_meal_allowance: float,
                    available_weekday_ot_hours: float | None,
                    available_weekend_ot_hours: float | None,
                    current_weekday_ot_hours: float = 0,
                    current_weekend_ot_hours: float = 0,
                    current_overtime_pay: float = 0, current_gross: float = 0,
                    current_employer_cost: float = 0, allow_taxable_bonus: bool = False,
                    current_performance_bonus: float = 0, pit_taxable_income: float = 0,
                    pit_zero_headroom: float = 0) -> dict[str, Any]:
    """Lập đề xuất thực lĩnh bảo thủ, chỉ dùng dư địa tiền ăn và giờ OT có thật."""
    if standard_days <= 0:
        raise ValueError("Ngày công chuẩn phải lớn hơn 0")
    supplied_hours = [value for value in (available_weekday_ot_hours,
                                           available_weekend_ot_hours) if value is not None]
    if supplied_hours and min(supplied_hours) < 0:
        raise ValueError("Số giờ làm thêm thực tế không được âm")
    values = (current_net, target_net, current_pit, current_employee_insurance,
              base_salary, current_meal_allowance, current_overtime_pay,
              current_gross, current_employer_cost, current_performance_bonus,
              pit_taxable_income, pit_zero_headroom)
    if min(values) < 0:
        raise ValueError("Dữ liệu tiền lương không được âm")

    meal_cap = 1_200_000 if month >= date(2026, 7, 1) else 730_000
    compliant_meal = min(current_meal_allowance, meal_cap)
    meal_compliance_delta = compliant_meal - current_meal_allowance
    remaining = max(_money(target_net) - _money(current_net + meal_compliance_delta), 0)
    meal_room = max(meal_cap - compliant_meal, 0)
    meal_increase = min(remaining, meal_room)
    remaining -= meal_increase
    meal_delta = meal_compliance_delta + meal_increase

    hourly_rate = base_salary / standard_days / 8
    weekday_rate = hourly_rate * 1.5
    weekend_rate = hourly_rate * 2
    current_total_hours = current_weekday_ot_hours + current_weekend_ot_hours
    proposed_weekday_base = current_weekday_ot_hours
    proposed_weekend_base = current_weekend_ot_hours
    overtime_compliance_delta = 0.0
    if available_weekday_ot_hours is None and available_weekend_ot_hours is None and current_total_hours > 40:
        proposed_weekday_base = min(current_weekday_ot_hours, 32 if current_weekend_ot_hours > 0 else 40)
        proposed_weekend_base = min(current_weekend_ot_hours, 40 - proposed_weekday_base)
        compliant_overtime_pay = _money(proposed_weekday_base * weekday_rate + proposed_weekend_base * weekend_rate)
        overtime_compliance_delta = compliant_overtime_pay - current_overtime_pay
        remaining = max(remaining - overtime_compliance_delta, 0)
    proposed_base_hours = proposed_weekday_base + proposed_weekend_base
    legal_hours_room = max(40 - proposed_base_hours, 0)
    weekday_capacity = legal_hours_room if available_weekday_ot_hours is None else min(available_weekday_ot_hours, legal_hours_room)
    weekend_capacity = max(legal_hours_room - weekday_capacity, 0)
    if available_weekend_ot_hours is not None:
        weekend_capacity = min(available_weekend_ot_hours, weekend_capacity)
    weekday_hours = min(weekday_capacity,
                        remaining / weekday_rate if weekday_rate else 0)
    weekday_increase = min(remaining, _money(weekday_hours * weekday_rate))
    remaining -= weekday_increase
    weekend_hours = min(weekend_capacity,
                        remaining / weekend_rate if weekend_rate else 0)
    weekend_increase = min(remaining, _money(weekend_hours * weekend_rate))
    remaining -= weekend_increase

    overtime_increase = weekday_increase + weekend_increase
    overtime_delta = overtime_compliance_delta + overtime_increase
    bonus_increase = 0.0
    proposed_pit = _money(current_pit)
    if allow_taxable_bonus and remaining > 0:
        def bonus_net(gross_bonus: int) -> float:
            taxable = pit_taxable_income + max(gross_bonus - pit_zero_headroom, 0)
            return gross_bonus - max(calculate_pit(taxable) - current_pit, 0)

        low, high = 0, int(max(remaining * 2, 1_000_000))
        while bonus_net(high) < remaining:
            high *= 2
        while low < high:
            middle = (low + high) // 2
            if bonus_net(middle) >= remaining:
                high = middle
            else:
                low = middle + 1
        bonus_increase = float(low)
        taxable = pit_taxable_income + max(bonus_increase - pit_zero_headroom, 0)
        proposed_pit = max(_money(current_pit), calculate_pit(taxable))

    pit_increase = max(proposed_pit - current_pit, 0)
    total_gross_delta = meal_delta + overtime_delta + bonus_increase
    total_net_delta = total_gross_delta - pit_increase
    proposed_net = _money(current_net + total_net_delta)
    shortfall = _money(max(target_net - proposed_net, 0))
    proposed_meal = _money(current_meal_allowance + meal_delta)
    total_weekday_hours = round(proposed_weekday_base + weekday_hours, 4)
    total_weekend_hours = round(proposed_weekend_base + weekend_hours, 4)
    cashflows = [
        {"key": "meal", "label": "Tiền ăn", "current": _money(current_meal_allowance),
         "proposed": proposed_meal, "delta": _signed_money(meal_delta)},
        {"key": "overtime", "label": "Tiền làm thêm hợp lệ", "current": _money(current_overtime_pay),
         "proposed": _money(current_overtime_pay + overtime_delta),
         "delta": _signed_money(overtime_delta)},
        {"key": "performance_bonus", "label": "Thưởng hiệu quả kinh doanh",
         "current": _money(current_performance_bonus),
         "proposed": _money(current_performance_bonus + bonus_increase),
         "delta": _money(bonus_increase)},
        {"key": "gross", "label": "Tổng thu nhập (gross)", "current": _money(current_gross),
         "proposed": _money(current_gross + total_gross_delta), "delta": _signed_money(total_gross_delta)},
        {"key": "pit", "label": "Thuế TNCN", "current": _money(current_pit),
         "proposed": _money(proposed_pit), "delta": _signed_money(pit_increase)},
        {"key": "insurance", "label": "BHXH người lao động",
         "current": _money(current_employee_insurance),
         "proposed": _money(current_employee_insurance), "delta": 0},
        {"key": "net", "label": "Thực lĩnh", "current": _money(current_net),
         "proposed": proposed_net, "delta": _signed_money(total_net_delta)},
        {"key": "employer_cost", "label": "Tổng chi phí công ty",
         "current": _money(current_employer_cost),
         "proposed": _money(current_employer_cost + total_gross_delta),
         "delta": _signed_money(total_gross_delta)},
    ]
    warnings = []
    if current_meal_allowance > meal_cap:
        cap_text = f"{int(meal_cap):,}".replace(",", ".")
        warnings.append(f"Tiền ăn hiện tại vượt trần; tự điều chỉnh về {cap_text} đồng.")
    if current_total_hours > 40:
        warnings.append(f"File đang có {current_total_hours:g} giờ làm thêm, vượt giới hạn 40 giờ/tháng; không đề xuất cộng thêm.")
    return {
        "feasible": shortfall == 0,
        "current_net": _money(current_net), "target_net": _money(target_net),
        "proposed_net": proposed_net, "shortfall": shortfall,
        "current_pit": _money(current_pit), "proposed_pit": _money(proposed_pit),
        "current_employee_insurance": _money(current_employee_insurance),
        "proposed_employee_insurance": _money(current_employee_insurance),
        "proposed": {
            "meal_allowance": proposed_meal,
            "overtime_weekday_hours": total_weekday_hours,
            "overtime_weekend_hours": total_weekend_hours,
            "attendance_bonus": _money(current_performance_bonus + bonus_increase) if allow_taxable_bonus else None,
            "performance_bonus": _money(current_performance_bonus + bonus_increase),
        },
        "cashflows": cashflows,
        "warnings": warnings,
        "dependencies": [
            "Tiền ăn phải được quy định trong hợp đồng lao động, thỏa ước hoặc quy chế công ty.",
            "Giờ làm thêm chỉ áp dụng theo bảng chấm công và phê duyệt làm thêm thực tế.",
            "Khoản hoàn chi điện thoại, xăng xe chỉ xem xét riêng khi có quy chế và chứng từ; hệ thống không tự cộng.",
            "Kế toán phải review bản nháp trước khi gửi bản Excel mới lên Google Drive.",
        ],
    }


def detect_month(*texts: str) -> str:
    combined = " ".join(texts).lower()
    match = re.search(r"(?:th[aá]ng\s*)?(1[0-2]|0?[1-9])\D+(20\d{2})", combined)
    if not match:
        match = re.search(r"(20\d{2})\D+(1[0-2]|0?[1-9])", combined)
        if match:
            return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
        return ""
    return f"{int(match.group(2)):04d}-{int(match.group(1)):02d}"


def review_workbook(path: Path) -> dict[str, Any]:
    """Doc bo cuc Excel cu va tra ve snapshot toi thieu, khong luu ten NV vao log."""
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    # Chế độ read-only tin vào dimension khai báo trong file; file tháng 1 khai báo
    # sai A1:A1000 nên phải mở bản giá trị bình thường để phục hồi đủ cột.
    value_workbook = load_workbook(path, data_only=True, read_only=False)
    value_sheet = value_workbook[value_workbook.sheetnames[0]]
    findings: list[dict[str, str]] = []
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    max_row = sheet.max_row or 0
    month = detect_month(sheet.title, path.name)
    rows = []
    employee_rows: list[int] = []
    for row in range(15, min(max_row, 200) + 1):
        code = sheet.cell(row, 2).value
        name = sheet.cell(row, 3).value
        if not code and not name:
            continue
        employee_rows.append(row)
        rows.append({"row": row, "code": str(code or ""), "has_name": bool(name)})
        meal = sheet.cell(row, COL_MEAL).value
        if month >= "2026-07" and isinstance(meal, (int, float)) and meal > 1_200_000:
            findings.append({"level": "vang", "code": "meal_cap",
                             "message": "Tiền ăn vượt 1,2 triệu đồng từ ngày 01/07/2026.",
                             "cells": [f"{COL_MEAL_LETTER}{row}"]})
    missing_kpcd = [row for row in employee_rows if sheet.cell(row, 19).value in (None, "")]
    if employee_rows and len(missing_kpcd) == len(employee_rows):
        findings.insert(0, {"level": "do", "code": "missing_kpcd",
                            "message": "Cột KPCĐ 2% đang trống ở các dòng nhân viên.",
                            "cells": [f"S{row}" for row in missing_kpcd]})
    grid: list[list[str]] = []
    # Một số file Excel khai báo sai dimension (ví dụ A1:A1000 dù dữ liệu có 40 cột).
    # Đọc theo tọa độ của sheet công thức để vẫn lấy đủ bảng, nhưng giới hạn 200x40.
    row_limit = min(max_row, 200)
    col_limit = min(max(sheet.max_column or 1, 40), 40)
    for row in islice(value_sheet.iter_rows(min_row=1, max_col=col_limit), row_limit):
        grid.append(["" if cell.value is None else str(cell.value) for cell in row])
    return {"sheet": sheet.title, "month": month, "rows": rows,
            "grid": grid, "ncols": max((len(row) for row in grid), default=0),
            "findings": findings}


def apply_workbook_changes(source: Path, output: Path, changes: list[dict[str, Any]]) -> None:
    """Tạo bản Excel mới, giữ nguyên bản Drive đã sync và công thức còn lại."""
    workbook = load_workbook(source, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    for change in changes:
        row = int(change["row"])
        if row < 15 or row > 200:
            raise ValueError("Dòng nhân viên không hợp lệ")
        if sheet.cell(row, 2).value in (None, "") and sheet.cell(row, 3).value in (None, ""):
            raise ValueError(f"Dòng {row} không có nhân viên")
        if change.get("meal_allowance") is not None:
            sheet.cell(row, COL_MEAL).value = float(change["meal_allowance"])
        if change.get("attendance_bonus") is not None:
            sheet.cell(row, 13).value = float(change["attendance_bonus"])
        if change.get("overtime_weekday_hours") is not None:
            hours = float(change["overtime_weekday_hours"])
            sheet.cell(row, 14).value = f"=(D{row}/E{row}/8)*1.5*{hours:g}"
        if change.get("overtime_weekend_hours") is not None:
            hours = float(change["overtime_weekend_hours"])
            sheet.cell(row, 15).value = f"=(D{row}/E{row}/8)*2*{hours:g}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
