"""Khu vực hợp tác riêng INUT - PYMID, giới hạn theo khách hàng."""
from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import audit, pymid
from .auth import CurrentUser, require_user
from .db import Customer, PymidCoopOrder, PymidCoopProduct, get_session

router = APIRouter(prefix="/api/pymid", tags=["pymid-coop"])
PYMID_TAX_CODE = "0313610275"


class OrderItemIn(BaseModel):
    product_id: int = Field(gt=0)
    quantity: float = Field(gt=0, le=100000)


class OrderIn(BaseModel):
    level: int = Field(ge=1, le=3)
    document_date: date = Field(default_factory=date.today)
    customer_reference: str = Field(default="", max_length=255)
    note: str = Field(default="", max_length=1000)
    items: list[OrderItemIn] = Field(min_length=1, max_length=100)


def _pymid_customer(db: Session) -> Customer:
    row = db.scalar(select(Customer).where(Customer.tax_code == PYMID_TAX_CODE))
    if not row:
        raise HTTPException(404, "Chưa có hồ sơ khách hàng CÔNG TY TNHH PYMID")
    return row


def _authorize(db: Session, user: CurrentUser) -> Customer:
    customer = _pymid_customer(db)
    if user.is_admin or user.customer_id == customer.id:
        return customer
    raise HTTPException(403, "Tài khoản không có quyền truy cập INUT - PYMID CO.OP")


def _loads(value: str) -> list:
    try:
        parsed = json.loads(value or "[]")
        return parsed if isinstance(parsed, list) else []
    except json.JSONDecodeError:
        return []


def _catalog_row(row: PymidCoopProduct, policy: dict) -> dict:
    pricing = pymid.price_for(row.source_price, row.category, policy)
    return {"id": row.id, "code": row.code, "name": row.name, "unit": row.unit,
            "source_price": round(row.source_price), "category": row.category,
            "level": row.level, "valid_from": row.valid_from, **pricing}


def _order_out(row: PymidCoopOrder) -> dict:
    return {"id": row.id, "customer_id": row.customer_id, "level": row.level,
            "status": row.status, "document_date": row.document_date,
            "customer_reference": row.customer_reference, "note": row.note,
            "policy_code": row.policy_code, "items": _loads(row.items_json),
            "invoice_lines": _loads(row.invoice_lines_json),
            "total_net": round(row.total_net), "total_tax": round(row.total_tax),
            "total_gross": round(row.total_gross),
            "created_at": row.created_at.isoformat(), "updated_at": row.updated_at.isoformat()}


def _build_snapshot(db: Session, body: OrderIn) -> tuple[dict, list[dict], list[dict]]:
    policy = pymid.vat_policy(body.document_date)
    ids = [item.product_id for item in body.items]
    products = {row.id: row for row in db.scalars(select(PymidCoopProduct).where(
        PymidCoopProduct.id.in_(ids), PymidCoopProduct.enabled.is_(True)))}
    if len(products) != len(set(ids)):
        raise HTTPException(400, "Có hạng mục không tồn tại hoặc đã ngừng sử dụng")
    snapshots = []
    for item in body.items:
        product = products[item.product_id]
        pricing = pymid.price_for(product.source_price, product.category, policy)
        quantity = float(item.quantity)
        snapshots.append({"product_id": product.id, "code": product.code, "name": product.name,
                          "unit": product.unit, "category": product.category, "quantity": quantity,
                          **pricing, "net_amount": pymid.money(pricing["net_price"] * quantity),
                          "tax_amount": pymid.money(pricing["tax_amount"] * quantity),
                          "gross_amount": pymid.money(pricing["gross_price"] * quantity)})
    hardware = [item for item in snapshots if item["category"] == "hardware"]
    software = [item for item in snapshots if item["category"] == "software"]
    lines = []
    if hardware:
        lines.append({"name": pymid.invoice_name(body.level), "unit": "Bộ", "quantity": 1,
                      "tax_treatment": "taxable", "vat_rate": policy["vat_rate"],
                      "vat_label": policy["label"],
                      "net_amount": sum(item["net_amount"] for item in hardware),
                      "tax_amount": sum(item["tax_amount"] for item in hardware),
                      "gross_amount": sum(item["gross_amount"] for item in hardware)})
    for item in software:
        lines.append({"name": f"iNut Nebi Software: License {item['name']}", "unit": "Gói",
                      "quantity": item["quantity"], "tax_treatment": "exempt", "vat_rate": None,
                      "vat_label": "KCT", "net_amount": item["net_amount"], "tax_amount": 0,
                      "gross_amount": item["gross_amount"]})
    return policy, snapshots, lines


@router.get("/catalog")
def catalog(document_date: date = date.today(), db: Session = Depends(get_session),
            user: CurrentUser = Depends(require_user)):
    _authorize(db, user)
    policy = pymid.vat_policy(document_date)
    rows = db.scalars(select(PymidCoopProduct).where(PymidCoopProduct.enabled.is_(True)).order_by(PymidCoopProduct.id))
    return {"policy": policy, "items": [_catalog_row(row, policy) for row in rows]}


@router.get("/orders")
def orders(db: Session = Depends(get_session), user: CurrentUser = Depends(require_user)):
    customer = _authorize(db, user)
    stmt = select(PymidCoopOrder).where(PymidCoopOrder.customer_id == customer.id).order_by(PymidCoopOrder.id.desc())
    return [_order_out(row) for row in db.scalars(stmt)]


@router.post("/orders")
def create_order(body: OrderIn, db: Session = Depends(get_session),
                 user: CurrentUser = Depends(require_user)):
    customer = _authorize(db, user)
    policy, items, lines = _build_snapshot(db, body)
    row = PymidCoopOrder(customer_id=customer.id, created_by=user.id, level=body.level,
                         document_date=body.document_date.isoformat(), customer_reference=body.customer_reference.strip(),
                         note=body.note.strip(), policy_code=policy["code"],
                         items_json=json.dumps(items, ensure_ascii=False),
                         invoice_lines_json=json.dumps(lines, ensure_ascii=False),
                         total_net=sum(item["net_amount"] for item in lines),
                         total_tax=sum(item["tax_amount"] for item in lines),
                         total_gross=sum(item["gross_amount"] for item in lines))
    db.add(row); db.commit(); db.refresh(row)
    audit.record(db, user.username, user.role, user.ip, "pymid_order_create", f"order:{row.id}",
                 f"Level {row.level} · {round(row.total_gross):,}đ")
    return _order_out(row)


def _get_order(db: Session, order_id: int, user: CurrentUser) -> PymidCoopOrder:
    customer = _authorize(db, user)
    row = db.get(PymidCoopOrder, order_id)
    if not row or row.customer_id != customer.id:
        raise HTTPException(404, "Không tìm thấy đơn hàng PYMID")
    return row


@router.post("/orders/{order_id}/submit")
def submit_order(order_id: int, db: Session = Depends(get_session),
                 user: CurrentUser = Depends(require_user)):
    row = _get_order(db, order_id, user)
    if row.status != "draft":
        raise HTTPException(409, "Chỉ đơn nháp mới được gửi duyệt")
    row.status = "submitted"; row.updated_at = datetime.now(); db.commit(); db.refresh(row)
    return _order_out(row)


@router.post("/orders/{order_id}/approve")
def approve_order(order_id: int, db: Session = Depends(get_session),
                  user: CurrentUser = Depends(require_user)):
    if not user.is_admin:
        raise HTTPException(403, "Chỉ INUT được duyệt đơn hàng")
    row = _get_order(db, order_id, user)
    if row.status not in {"draft", "submitted"}:
        raise HTTPException(409, "Đơn hàng không ở trạng thái có thể duyệt")
    row.status = "approved"; row.updated_at = datetime.now(); db.commit(); db.refresh(row)
    audit.record(db, user.username, user.role, user.ip, "pymid_order_approve", f"order:{row.id}", "Đã duyệt")
    return _order_out(row)


def _style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="1F5A4A")
    for cell in ws[1]:
        cell.fill = fill; cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def _order_workbook(row: PymidCoopOrder) -> bytes:
    wb = Workbook(); quote = wb.active; quote.title = "Báo giá"
    quote.append(["Tên hạng mục", "ĐVT", "SL", "Tiền trước thuế", "Thuế", "VAT", "Thanh toán"])
    for line in _loads(row.invoice_lines_json):
        quote.append([line["name"], line["unit"], line["quantity"], line["net_amount"],
                      line["tax_amount"], line["vat_label"], line["gross_amount"]])
    quote.append(["TỔNG", "", "", round(row.total_net), round(row.total_tax), "", round(row.total_gross)])
    _style_sheet(quote)

    nhanh = wb.create_sheet("Nhanh.vn")
    nhanh.append(["Tên sản phẩm", "Mã sản phẩm", "Đơn vị tính", "Số lượng", "Giá bán", "VAT", "Ghi chú"])
    for index, line in enumerate(_loads(row.invoice_lines_json), 1):
        nhanh.append([line["name"], f"NEBI-L{row.level}-{index}", line["unit"], line["quantity"],
                      line["net_amount"], line["vat_label"], row.customer_reference])
    _style_sheet(nhanh)

    config = wb.create_sheet("Cấu hình")
    config.append(["Mã", "Hạng mục", "Loại", "ĐVT", "SL", "Đơn giá thanh toán", "Thành tiền"])
    for item in _loads(row.items_json):
        config.append([item["code"], item["name"], item["category"], item["unit"], item["quantity"],
                       item["gross_price"], item["gross_amount"]])
    _style_sheet(config)
    output = BytesIO(); wb.save(output); return output.getvalue()


@router.get("/orders/{order_id}/xlsx")
def export_order(order_id: int, db: Session = Depends(get_session),
                 user: CurrentUser = Depends(require_user)):
    row = _get_order(db, order_id, user)
    filename = f"INUT-PYMID-COOP-Level-{row.level}-{row.id}.xlsx"
    return Response(_order_workbook(row), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})
