from __future__ import annotations

from datetime import date

import pytest
from fastapi.testclient import TestClient

from openpyxl import Workbook, load_workbook

from app.payroll import (
    COL_MEAL,
    PayrollInput,
    apply_workbook_changes,
    calculate_payroll,
    calculate_pit,
    insurance_base_cap,
    plan_net_target,
    review_workbook,
)
from app.payroll_api import _formula_hours


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NAS_ENABLED", "false")
    from app.config import get_settings
    from app import db as dbmod
    from app.auth import ensure_admin_seed

    get_settings.cache_clear()
    dbmod.reset_engine_for_tests()
    dbmod.init_db()
    gen = dbmod.get_session()
    session = next(gen)
    ensure_admin_seed(session, get_settings())
    gen.close()
    from app.main import app

    return TestClient(app)


def _login(client):
    response = client.post(
        "/api/login", json={"username": "admin", "password": "NhapHang123@"}
    )
    assert response.status_code == 200


def test_pit_2026_uses_five_bands_and_new_family_deductions():
    assert calculate_pit(10_000_000) == 500_000
    assert calculate_pit(30_000_000) == 2_500_000
    assert calculate_pit(60_000_000) == 8_500_000
    assert calculate_pit(100_000_000) == 20_500_000
    assert calculate_pit(110_000_000) == 24_000_000


def test_july_meal_allowance_caps_exempt_amount_and_overtime_is_exempt():
    result = calculate_payroll(
        PayrollInput(
            month=date(2026, 7, 1),
            base_salary=20_000_000,
            actual_days=22,
            standard_days=22,
            meal_allowance=1_500_000,
            overtime_pay=2_000_000,
        )
    )
    assert result.gross_income == 23_500_000
    assert result.taxable_income_before_deductions == 20_300_000
    assert result.employee_insurance == 2_100_000
    assert result.pit_taxable_income == 2_700_000
    assert result.pit == 135_000


def test_meal_allowance_before_july_uses_730k_cap():
    result = calculate_payroll(PayrollInput(
        month=date(2026, 6, 1), base_salary=20_000_000, standard_days=22,
        actual_days=22, meal_allowance=1_000_000,
    ))
    assert result.taxable_income_before_deductions == 20_270_000


def test_net_target_uses_lawful_meal_room_then_actual_overtime_without_more_tax_or_insurance():
    plan = plan_net_target(
        month=date(2026, 7, 1), current_net=15_000_000, target_net=17_300_000,
        current_pit=350_000, current_employee_insurance=1_260_000,
        base_salary=12_000_000, standard_days=22, current_meal_allowance=500_000,
        available_weekday_ot_hours=16, available_weekend_ot_hours=8,
    )

    assert plan["feasible"] is True
    assert plan["proposed"]["meal_allowance"] == 1_200_000
    assert plan["proposed"]["overtime_weekday_hours"] > 0
    assert plan["proposed"]["overtime_weekend_hours"] == 0
    assert plan["proposed_net"] == 17_300_000
    assert plan["proposed_pit"] == plan["current_pit"] == 350_000
    assert plan["proposed_employee_insurance"] == plan["current_employee_insurance"] == 1_260_000
    assert any(row["key"] == "net" and row["delta"] == 2_300_000 for row in plan["cashflows"])


def test_net_target_reports_shortfall_and_does_not_invent_unproven_allowances():
    plan = plan_net_target(
        month=date(2026, 6, 1), current_net=10_000_000, target_net=20_000_000,
        current_pit=0, current_employee_insurance=1_050_000,
        base_salary=10_000_000, standard_days=22, current_meal_allowance=700_000,
        available_weekday_ot_hours=0, available_weekend_ot_hours=0,
    )

    assert plan["feasible"] is False
    assert plan["proposed"]["meal_allowance"] == 730_000
    assert plan["shortfall"] == 9_970_000
    assert plan["proposed"]["attendance_bonus"] is None
    assert "chứng từ" in " ".join(plan["dependencies"]).lower()


def test_net_target_rejects_invalid_days_and_negative_actual_overtime():
    with pytest.raises(ValueError, match="Ngày công chuẩn"):
        plan_net_target(
            month=date(2026, 7, 1), current_net=1, target_net=2, current_pit=0,
            current_employee_insurance=0, base_salary=1, standard_days=0,
            current_meal_allowance=0, available_weekday_ot_hours=0,
            available_weekend_ot_hours=0,
        )


def test_net_target_automatically_normalizes_meal_and_respects_monthly_overtime_limit():
    plan = plan_net_target(
        month=date(2026, 7, 1), current_net=28_687_885, target_net=33_345_678,
        current_pit=0, current_employee_insurance=577_500,
        base_salary=5_500_000, standard_days=26, current_meal_allowance=1_500_000,
        available_weekday_ot_hours=None, available_weekend_ot_hours=None,
        current_weekday_ot_hours=128, current_weekend_ot_hours=48,
        current_overtime_pay=7_615_385, current_gross=29_265_385,
        current_employer_cost=30_447_885,
    )

    assert plan["feasible"] is False
    assert plan["proposed"]["meal_allowance"] == 1_200_000
    assert plan["proposed"]["overtime_weekday_hours"] == 128
    assert plan["proposed"]["overtime_weekend_hours"] == 48
    assert plan["proposed_net"] == 28_387_885
    assert plan["shortfall"] == 4_957_793
    assert any("176" in warning and "40" in warning for warning in plan["warnings"])
    with pytest.raises(ValueError, match="làm thêm"):
        plan_net_target(
            month=date(2026, 7, 1), current_net=1, target_net=2, current_pit=0,
            current_employee_insurance=0, base_salary=1, standard_days=22,
            current_meal_allowance=0, available_weekday_ot_hours=-1,
            available_weekend_ot_hours=0,
        )


def test_employer_cost_includes_insurance_and_trade_union_fee():
    result = calculate_payroll(
        PayrollInput(
            month=date(2026, 6, 1),
            base_salary=10_000_000,
            actual_days=22,
            standard_days=22,
            insurance_salary=10_000_000,
        )
    )
    assert result.employer_insurance == 2_150_000
    assert result.trade_union_fee == 200_000
    assert result.total_employer_cost == 12_350_000


def test_review_legacy_workbook_flags_missing_kpcd_and_july_meal_cap(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Thang 7-2026"
    sheet.cell(15, 2, "NV-DEMO")
    sheet.cell(15, 3, "Nhan vien mau")
    sheet.cell(15, COL_MEAL, 1_500_000)  # cot I: tien an
    path = tmp_path / "bang-luong-an-danh.xlsx"
    workbook.save(path)

    review = review_workbook(path)
    assert review["rows"] == [{"row": 15, "code": "NV-DEMO", "has_name": True}]
    assert {finding["code"] for finding in review["findings"]} == {"missing_kpcd", "meal_cap"}

    sheet.title = "Thang 8-2026"
    workbook.save(path)
    august = review_workbook(path)
    assert august["month"] == "2026-08"
    assert any(finding["code"] == "meal_cap" for finding in august["findings"])


def test_review_empty_workbook_does_not_crash(tmp_path):
    workbook = Workbook()
    path = tmp_path / "empty.xlsx"
    workbook.save(path)
    review = review_workbook(path)
    assert review["rows"] == []
    assert review["findings"] == []


def test_insurance_base_is_capped_at_twenty_times_reference():
    # Dieu 31 Luat BHXH 2024: tran = 20 x muc tham chieu (luong co so).
    # Truoc 01/07/2026 luong co so 2,34tr -> tran 46,8tr; tu 01/07/2026 2,53tr -> 50,6tr.
    assert insurance_base_cap(date(2026, 6, 1)) == 46_800_000
    assert insurance_base_cap(date(2026, 7, 1)) == 50_600_000

    # Luong cao hon tran -> chi dong tren phan tran, khong dong tren toan bo luong.
    high = calculate_payroll(PayrollInput(
        month=date(2026, 7, 1), base_salary=80_000_000, standard_days=22, actual_days=22,
    ))
    assert high.insurance_salary == 50_600_000
    assert high.employee_insurance == 5_313_000  # 50,6tr x 10,5%
    assert high.employer_insurance == 10_879_000  # 50,6tr x 21,5%

    # Luong duoi tran -> giu nguyen cach tinh cu (khong hoi quy).
    low = calculate_payroll(PayrollInput(
        month=date(2026, 7, 1), base_salary=20_000_000, standard_days=22, actual_days=22,
    ))
    assert low.insurance_salary == 20_000_000
    assert low.employee_insurance == 2_100_000


def test_formula_hours_reads_hours_not_rate_times_hours():
    # Cong thuc do apply_workbook_changes sinh ra: gio phai la factor cuoi, KHONG nhan 1.5/2.
    assert _formula_hours("=(D15/E15/8)*1.5*12") == 12
    assert _formula_hours("=(D15/E15/8)*2*8") == 8
    assert _formula_hours("=(D15/E15/8)*1.5*0") == 0
    assert _formula_hours("=(D15/E15/8)*1.5*(8*16)") == 128
    assert _formula_hours("=(D15/E15/8)*2*(8*6)") == 48
    # Cong thuc khong co phan gio -> khong bat nham so 8 trong (D/E/8).
    assert _formula_hours("=D15/E15/8") == 0
    assert _formula_hours(12000) == 0
    assert _formula_hours(None) == 0


def test_review_and_edit_agree_on_meal_column(tmp_path):
    # Mot fixture DUNG CHUNG cho ca review va edit: bat lech cot tien an neu tai dien.
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Thang 7-2026"
    sheet.cell(15, 2, "NV-DEMO")
    sheet.cell(15, 3, "Nhan vien mau")
    sheet.cell(15, COL_MEAL, 1_500_000)  # tien an vuot tran o cot I
    sheet.cell(15, 19, 240_000)  # cot S: KPCD da co -> khong bao missing_kpcd
    source = tmp_path / "source.xlsx"
    workbook.save(source)

    # Review phai bat duoc canh bao vuot tran o dung cot I.
    before = review_workbook(source)
    meal_findings = [f for f in before["findings"] if f["code"] == "meal_cap"]
    assert meal_findings, "review phai bat tien an vuot tran o cot I"
    assert meal_findings[0]["cells"] == ["I15"]

    # Edit ha tien an ve muc hop le -> re-review phai het canh bao meal_cap.
    output = tmp_path / "draft.xlsx"
    apply_workbook_changes(source, output, [{"row": 15, "meal_allowance": 1_200_000}])
    after = review_workbook(output)
    assert not any(f["code"] == "meal_cap" for f in after["findings"])


def test_payroll_period_workflow_and_override_reason(client):
    _login(client)
    employee = client.post(
        "/api/payroll/employees",
        json={
            "code": "NV-TEST-01",
            "name": "Nhan vien A",
            "base_salary": 12_000_000,
            "insurance_salary": 12_000_000,
            "dependents": 0,
        },
    )
    assert employee.status_code == 200, employee.text

    period = client.post("/api/payroll/periods", json={"month": "2026-07"})
    assert period.status_code == 200, period.text
    period_id = period.json()["id"]
    line_id = period.json()["lines"][0]["id"]

    no_reason = client.put(
        f"/api/payroll/periods/{period_id}/lines/{line_id}",
        json={"override_net": 99_000_000},
    )
    assert no_reason.status_code == 422

    updated = client.put(
        f"/api/payroll/periods/{period_id}/lines/{line_id}",
        json={"actual_days": 22, "standard_days": 22},
    )
    assert updated.status_code == 200, updated.text

    reviewed = client.post(f"/api/payroll/periods/{period_id}/review")
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["status"] == "reviewed"

    locked = client.post(f"/api/payroll/periods/{period_id}/lock")
    assert locked.status_code == 200, locked.text
    assert locked.json()["status"] == "locked"

    rejected = client.put(
        f"/api/payroll/periods/{period_id}/lines/{line_id}",
        json={"actual_days": 21, "standard_days": 22},
    )
    assert rejected.status_code == 409


def test_drive_sync_runs_as_job_and_reports_progress(client, monkeypatch):
    from app.config import get_settings
    from app import payroll_api

    target = get_settings().data_path / "payroll_drive"
    target.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Thang 7-2026"
    sheet.cell(15, 2, "NV-DEMO")
    workbook.save(target / "payroll-demo.xlsx")
    monkeypatch.setattr(
        payroll_api.subprocess,
        "run",
        lambda *args, **kwargs: payroll_api.subprocess.CompletedProcess(args[0], 0, "", ""),
    )

    _login(client)
    started = client.post("/api/payroll/sync-drive")
    assert started.status_code == 200, started.text
    assert started.json()["job_id"] > 0

    status = client.get("/api/payroll/sync-drive/status")
    assert status.status_code == 200
    job = status.json()["job"]
    assert job["status"] == "success"
    assert job["stats"]["progress"] == 100
    assert job["stats"]["files"] == ["payroll-demo.xlsx"]

    imports = client.get("/api/payroll/imports")
    assert imports.status_code == 200
    assert imports.json()[0]["filename"] == "payroll-demo.xlsx"
    detail = client.get(f"/api/payroll/imports/{imports.json()[0]['id']}")
    assert detail.status_code == 200
    assert detail.json()["snapshot"]["grid"][14][1] == "NV-DEMO"


def test_apply_workbook_changes_updates_meal_bonus_and_overtime_formulas(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tháng 7"
    sheet["B15"] = "NV-DEMO"
    sheet["C15"] = "Nhân viên mẫu"
    sheet["D15"] = 12_000_000
    sheet["E15"] = 26
    source = tmp_path / "source.xlsx"
    output = tmp_path / "draft.xlsx"
    workbook.save(source)

    apply_workbook_changes(source, output, [{
        "row": 15,
        "meal_allowance": 1_200_000,
        "attendance_bonus": 2_000_000,
        "overtime_weekday_hours": 12,
        "overtime_weekend_hours": 8,
    }])

    changed = load_workbook(output, data_only=False).active
    assert changed["I15"].value == 1_200_000
    assert changed["M15"].value == 2_000_000
    assert changed["N15"].value == "=(D15/E15/8)*1.5*12"
    assert changed["O15"].value == "=(D15/E15/8)*2*8"


def test_import_draft_workflow_save_review_and_upload(client, monkeypatch):
    from app.config import get_settings
    from app import payroll_api

    target = get_settings().data_path / "payroll_drive"
    target.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tháng 7"
    sheet["B15"] = "NV-DEMO"
    sheet["C15"] = "Nhân viên mẫu"
    sheet["D15"] = 12_000_000
    sheet["E15"] = 26
    sheet["I15"] = 1_500_000
    sheet["S15"] = 240_000
    workbook.save(target / "payroll-edit-demo.xlsx")
    commands = []
    def successful_rclone(command, **kwargs):
        commands.append(command)
        return payroll_api.subprocess.CompletedProcess(command, 0, "", "")
    monkeypatch.setattr(payroll_api.subprocess, "run", successful_rclone)
    _login(client)
    client.post("/api/payroll/sync-drive")
    imported = next(row for row in client.get("/api/payroll/imports").json()
                    if row["filename"] == "payroll-edit-demo.xlsx")

    saved = client.post(f"/api/payroll/imports/{imported['id']}/draft", json={"changes": [{
        "row": 15, "meal_allowance": 1_200_000, "attendance_bonus": 2_000_000,
        "overtime_weekday_hours": 12, "overtime_weekend_hours": 8,
        "reason": "Tháng có nhiều hợp đồng",
    }]})
    assert saved.status_code == 200, saved.text
    draft_id = saved.json()["id"]
    assert saved.json()["status"] == "draft"

    reviewed = client.post(f"/api/payroll/drafts/{draft_id}/review")
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["status"] == "reviewed"
    assert not any(f["code"] == "meal_cap" for f in reviewed.json()["findings"])

    uploaded = client.post(f"/api/payroll/drafts/{draft_id}/upload-drive")
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()["status"] == "uploaded"
    assert uploaded.json()["drive_filename"] == "payroll-edit-demo.xlsx"
    upload_commands = [command for command in commands if "copyto" in command]
    assert len(upload_commands) == 2
    assert "bản gốc trước cập nhật" in upload_commands[0][3]
    assert upload_commands[1][3] == "vnmap-drive:payroll-edit-demo.xlsx"


def test_import_net_target_returns_comparison_without_changing_workbook(client, monkeypatch):
    from app.config import get_settings
    from app import payroll_api

    target = get_settings().data_path / "payroll_drive"
    target.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tháng 7-2026"
    values = {"B15": "NV-DEMO", "C15": "Nhân viên mẫu", "D15": 12_000_000,
              "E15": 22, "I15": 500_000, "Q15": 13_000_000, "R15": 12_000_000,
              "AG15": 350_000, "AJ15": 11_390_000, "AK15": 15_820_000}
    for cell, value in values.items():
        sheet[cell] = value
    sheet["N15"] = "=(D15/E15/8)*1.5*4"
    workbook.save(target / "payroll-target-demo.xlsx")
    monkeypatch.setattr(
        payroll_api.subprocess, "run",
        lambda *args, **kwargs: payroll_api.subprocess.CompletedProcess(args[0], 0, "", ""),
    )
    _login(client)
    client.post("/api/payroll/sync-drive")
    imported = next(row for row in client.get("/api/payroll/imports").json()
                    if row["filename"] == "payroll-target-demo.xlsx")

    response = client.post(f"/api/payroll/imports/{imported['id']}/net-target", json={
        "row": 15, "target_net": 13_000_000,
    })

    assert response.status_code == 200, response.text
    assert response.json()["current_net"] == 11_390_000
    assert response.json()["proposed_pit"] == 350_000
    assert response.json()["proposed_employee_insurance"] == 1_260_000
    assert response.json()["proposed"]["meal_allowance"] == 1_200_000
    assert response.json()["proposed"]["overtime_weekday_hours"] > 4


def test_import_net_target_validates_employee_row(client):
    _login(client)
    response = client.post("/api/payroll/imports/999/net-target", json={
        "row": 15, "target_net": 13_000_000,
        "available_weekday_ot_hours": 0, "available_weekend_ot_hours": 0,
    })
    assert response.status_code == 404
