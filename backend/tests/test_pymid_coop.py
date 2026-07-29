from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture
def app_env(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NAS_ENABLED", "false")
    from app.config import get_settings
    from app import db as dbmod
    from app.auth import ensure_admin_seed
    from app.security import hash_password

    get_settings.cache_clear()
    dbmod.reset_engine_for_tests()
    dbmod.init_db()
    gen = dbmod.get_session(); db = next(gen)
    ensure_admin_seed(db, get_settings())
    pymid = dbmod.Customer(name="CÔNG TY TNHH PYMID", tax_code="0313610275")
    other = dbmod.Customer(name="CÔNG TY KHÁC", tax_code="0100000000")
    db.add_all([pymid, other]); db.flush()
    db.add_all([
        dbmod.User(username="pymid", password_hash=hash_password("PymidTest123@"), role="customer", customer_id=pymid.id),
        dbmod.User(username="other", password_hash=hash_password("OtherTest123@"), role="customer", customer_id=other.id),
    ])
    db.commit(); gen.close()
    from app.main import app
    return TestClient(app)


def login(client: TestClient, username: str, password: str):
    result = client.post("/api/login", json={"username": username, "password": password})
    assert result.status_code == 200, result.text


def test_catalog_prices_and_customer_scope(app_env):
    login(app_env, "pymid", "PymidTest123@")
    response = app_env.get("/api/pymid/catalog?document_date=2026-07-29")
    assert response.status_code == 200, response.text
    rows = response.json()["items"]
    assert len(rows) == 23
    pmc01 = next(row for row in rows if row["code"] == "PMC01")
    assert pmc01["gross_price"] == 2_823_250
    assert pmc01["vat_rate"] == 8
    software = next(row for row in rows if row["code"] == "SW-PHUN-SUONG")
    assert software["gross_price"] == 632_500
    assert software["tax_treatment"] == "exempt"
    assert software["vat_label"] == "KCT"

    other = TestClient(app_env.app)
    login(other, "other", "OtherTest123@")
    assert other.get("/api/pymid/catalog").status_code == 403


def test_order_aggregates_hardware_and_keeps_software_separate(app_env):
    login(app_env, "pymid", "PymidTest123@")
    catalog = app_env.get("/api/pymid/catalog?document_date=2026-07-29").json()["items"]
    pmc01 = next(row for row in catalog if row["code"] == "PMC01")
    relay = next(row for row in catalog if row["code"] == "RS485-RELAY4")
    software = next(row for row in catalog if row["code"] == "SW-PHUN-SUONG")
    created = app_env.post("/api/pymid/orders", json={
        "level": 1,
        "document_date": "2026-07-29",
        "customer_reference": "Nhà yến mẫu",
        "note": "Cấu hình thử",
        "items": [
            {"product_id": pmc01["id"], "quantity": 1},
            {"product_id": relay["id"], "quantity": 2},
            {"product_id": software["id"], "quantity": 1},
        ],
    })
    assert created.status_code == 200, created.text
    order = created.json()
    assert order["status"] == "draft"
    assert len(order["invoice_lines"]) == 2
    hardware = next(line for line in order["invoice_lines"] if line["tax_treatment"] == "taxable")
    assert hardware["name"].startswith("iNut Nebi - Bộ giải pháp nhà yến")
    assert hardware["unit"] == "Bộ" and hardware["quantity"] == 1
    assert hardware["vat_rate"] == 8
    assert hardware["gross_amount"] == 2_823_250 + 2 * 632_500
    sw_line = next(line for line in order["invoice_lines"] if line["tax_treatment"] == "exempt")
    assert sw_line["vat_label"] == "KCT"
    assert sw_line["gross_amount"] == 632_500

    exported = app_env.get(f"/api/pymid/orders/{order['id']}/xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), data_only=True)
    assert {"Báo giá", "Nhanh.vn", "Cấu hình"}.issubset(workbook.sheetnames)
    assert workbook["Nhanh.vn"]["A2"].value.startswith("iNut Nebi")


def test_vat_policy_uses_10_percent_after_reduction_period(app_env):
    login(app_env, "admin", "NhapHang123@")
    response = app_env.get("/api/pymid/catalog?document_date=2027-01-01")
    assert response.status_code == 200
    pmc01 = next(row for row in response.json()["items"] if row["code"] == "PMC01")
    assert pmc01["vat_rate"] == 10
    assert pmc01["gross_price"] == 2_875_532
    assert response.json()["policy"]["code"] == "VAT_STANDARD_10"


def test_only_admin_can_approve_order(app_env):
    login(app_env, "pymid", "PymidTest123@")
    pmc01 = next(row for row in app_env.get("/api/pymid/catalog").json()["items"] if row["code"] == "PMC01")
    order = app_env.post("/api/pymid/orders", json={
        "level": 1, "items": [{"product_id": pmc01["id"], "quantity": 1}],
    }).json()
    assert app_env.post(f"/api/pymid/orders/{order['id']}/approve").status_code == 403
    assert app_env.post(f"/api/pymid/orders/{order['id']}/submit").json()["status"] == "submitted"

    admin = TestClient(app_env.app)
    login(admin, "admin", "NhapHang123@")
    approved = admin.post(f"/api/pymid/orders/{order['id']}/approve")
    assert approved.status_code == 200
    assert approved.json()["status"] == "approved"


def test_draft_can_be_updated_but_submitted_order_is_locked(app_env):
    login(app_env, "pymid", "PymidTest123@")
    rows = app_env.get("/api/pymid/catalog?document_date=2026-07-29").json()["items"]
    base = next(row for row in rows if row["code"] == "PMC01")
    relay = next(row for row in rows if row["code"] == "RS485-RELAY4")
    created = app_env.post("/api/pymid/orders", json={
        "level": 1,
        "document_date": "2026-07-29",
        "customer_reference": "Bản nháp đầu",
        "items": [{"product_id": base["id"], "quantity": 1}],
    }).json()

    updated = app_env.put(f"/api/pymid/orders/{created['id']}", json={
        "level": 1,
        "document_date": "2026-07-29",
        "customer_reference": "Nhà yến đã sửa",
        "note": "Thêm relay",
        "items": [
            {"product_id": base["id"], "quantity": 1},
            {"product_id": relay["id"], "quantity": 2},
        ],
    })
    assert updated.status_code == 200, updated.text
    assert updated.json()["customer_reference"] == "Nhà yến đã sửa"
    assert len(updated.json()["items"]) == 2

    app_env.post(f"/api/pymid/orders/{created['id']}/submit")
    locked = app_env.put(f"/api/pymid/orders/{created['id']}", json={
        "level": 1,
        "items": [{"product_id": base["id"], "quantity": 1}],
    })
    assert locked.status_code == 409
    assert locked.json()["detail"] == "Đơn đã gửi duyệt nên không thể sửa"


def test_catalog_excel_is_available_to_pymid(app_env):
    login(app_env, "pymid", "PymidTest123@")
    exported = app_env.get("/api/pymid/catalog.xlsx?document_date=2026-07-29")
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content), data_only=True)
    assert {"Danh mục giá", "Nhanh.vn"}.issubset(workbook.sheetnames)
    assert workbook["Danh mục giá"].max_row == 24
    assert workbook["Danh mục giá"]["J2"].value == "VAT 8%"


def test_order_rejects_mismatched_level_and_duplicate_products(app_env):
    login(app_env, "pymid", "PymidTest123@")
    rows = app_env.get("/api/pymid/catalog").json()["items"]
    level_one = next(row for row in rows if row["code"] == "PMC01")

    mismatched = app_env.post("/api/pymid/orders", json={
        "level": 2,
        "items": [{"product_id": level_one["id"], "quantity": 1}],
    })
    assert mismatched.status_code == 400
    assert mismatched.json()["detail"] == "Cấu hình phải có đúng bộ trung tâm của Level đã chọn"

    duplicated = app_env.post("/api/pymid/orders", json={
        "level": 1,
        "items": [
            {"product_id": level_one["id"], "quantity": 1},
            {"product_id": level_one["id"], "quantity": 1},
        ],
    })
    assert duplicated.status_code == 400
    assert duplicated.json()["detail"] == "Mỗi hạng mục chỉ được xuất hiện một lần"


def test_pymid_owner_can_create_staff_account_with_coop_only_access(app_env):
    login(app_env, "pymid", "PymidTest123@")
    created = app_env.post("/api/pymid/staff", json={
        "username": "pymid.nhanvien01",
        "display_name": "Nhân viên kinh doanh 01",
        "password": "PymidStaff123@",
    })
    assert created.status_code == 201, created.text
    assert created.json() == {
        "id": created.json()["id"],
        "username": "pymid.nhanvien01",
        "display_name": "Nhân viên kinh doanh 01",
        "role": "pymid_staff",
    }

    staff_rows = app_env.get("/api/pymid/staff")
    assert staff_rows.status_code == 200
    assert [row["username"] for row in staff_rows.json()] == ["pymid.nhanvien01"]

    staff = TestClient(app_env.app)
    login(staff, "pymid.nhanvien01", "PymidStaff123@")
    assert staff.get("/api/pymid/catalog").status_code == 200
    me = staff.get("/api/me").json()
    assert me["role"] == "pymid_staff"
    assert me["portal_scope"] == "pymid_coop"
    assert staff.get("/api/my/documents").status_code == 403
    assert staff.get("/api/my/invoices").status_code == 403
    assert staff.get("/api/my/download.zip").status_code == 403
    assert staff.get("/api/pymid/staff").status_code == 403


def test_non_pymid_customer_cannot_create_pymid_staff(app_env):
    login(app_env, "other", "OtherTest123@")
    response = app_env.post("/api/pymid/staff", json={
        "username": "khongduocphep",
        "display_name": "Không được phép",
        "password": "OtherStaff123@",
    })
    assert response.status_code == 403
